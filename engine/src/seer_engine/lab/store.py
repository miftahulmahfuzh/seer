"""The method lab's database: ``lab/lab.sqlite``, committed, append-only (method lab design §1).

Impure edge: stdlib ``sqlite3`` and the clock (timestamps). Every lab read and write goes
through this module.

Tables:

- ``methods``: one row per idea. ``H-*`` ids are the historical record (P7a families, A, A2,
  B); ``M0001``… are lab methods. ``status`` only moves along ``TRANSITIONS`` (a trigger refuses
  any other update); ``analysis`` only grows (a trigger refuses a rewrite); ``hypothesis`` is
  frozen once the method leaves ``idea``; ``source_sha`` (the method file's sha256 when it ran)
  is set once.
- ``trials``: one row per backtest, the multiple-testing count. Triggers refuse every UPDATE
  and DELETE. ``UNIQUE(config_digest, window)``: a configuration runs at most once on the dev
  window and gets at most one look at the test window.
- ``ideas_seen``: dedupe keys (``url:…``, ``concept:…``) of sources already explored.
- ``insights``: the lab journal (observations, hypotheses, data and feature wishes, risks, and
  batch syntheses). Triggers refuse every UPDATE and DELETE.
- ``trial_moments``: the deflated Sharpe's inputs for one dev trial -- its daily Sharpe, the
  number of daily returns, the skew, the kurtosis, and the ``var_trials`` and N it was judged
  against -- so a recorded verdict can be recomputed later without re-running the backtest.
  At most one row per trial, keyed by ``trials.n``. Triggers refuse every UPDATE and DELETE.

Schema versions (``meta.schema_version``): 1 is the first lab; 2 adds the ``synthesis`` insight
kind; 3 adds the ``trial_moments`` side table. ``connect`` migrates an older database in place;
``connect_readonly`` never does.

The **verdict** a trial reads is derived, not frozen: ``DSR_MIN`` is the threshold (0.90 since
2026-10-07) and ``DSR_POLICY`` names the multiple-testing N (``all-trials`` today, so N is every
dev trial, as it has always been). ``verdict`` re-decides every condition at call time -- the four
threshold owner conditions from the trial's own recorded columns against the live constants, the
luck test on the trial's DSR at the gate's current N -- and carries only ``owner inputs`` from the
record, because no constant re-decides it. Recorded rows keep the labels of the bars they were
judged under, so every reader uses ``is_luck_label`` rather than comparing to ``DSR_LABEL``, and
``owner_failures`` rather than parsing a ``failed`` string.

``snapshot`` / ``snapshot_json`` turn a database (v1, v2 or v3) into the web's ``web/data/lab.json``
(seertrade.site/sera). ``lab stage`` writes it next to the database and stages both.

``journal_mode`` stays DELETE, so the committed file is the whole database (no ``-wal``).
"""

from __future__ import annotations

import json
import math
import os
import sqlite3
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, fields
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from seer_engine import config

COMMITTED_DB = config.REPO_ROOT / "lab" / "lab.sqlite"
# Parallel explorer sessions (sera-the-explorer) run in worktrees and share the main checkout's
# database through SEER_LAB_DB; only the coordinator commits it.
DB_PATH = Path(os.environ.get("SEER_LAB_DB") or COMMITTED_DB)
XLSX_PATH = config.REPO_ROOT / "lab" / "lab.xlsx"  # gitignored
SCHEMA_VERSION = "3"  # 2: the synthesis insight kind; 3: the trial_moments side table (see _migrate)

SOURCE_KINDS: tuple[str, ...] = ("paper", "blog", "github", "knowledge", "variation", "seed")
STATUSES: tuple[str, ...] = (
    "idea",
    "registered",
    "rejected",
    "dev-eligible",
    "promoted",
    "test-passed",
    "test-failed",
    "paper",
    "blocked-data",
)
TRANSITIONS: tuple[tuple[str, str], ...] = (
    ("idea", "registered"),
    ("idea", "rejected"),  # dropped before running (duplicate, untestable for another reason)
    ("idea", "blocked-data"),
    ("registered", "rejected"),
    ("registered", "dev-eligible"),
    ("registered", "blocked-data"),
    ("blocked-data", "idea"),  # the missing data arrived
    # The one edge out of ``rejected`` (LAB_LUCK_GATE_PLAN.md Decision D2). "Status moves
    # forward only" protects *verdicts from being erased*, and nothing here erases one: the
    # trials that decided the rejection stay in the append-only ``trials`` table, untouched, and
    # the re-evaluation is *appended* to the method's analysis. Admitted only by
    # ``reevaluate_method``, and only for a method whose dev trial clears all five conditions at
    # the bars in force today. Without this edge R1 is unsatisfiable: M0022 reads ``rejected``
    # today, and re-judging it under a new id would need a duplicate configuration digest, which
    # ``has_trial`` refuses outright.
    ("rejected", "dev-eligible"),
    ("dev-eligible", "promoted"),
    ("promoted", "test-passed"),
    ("promoted", "test-failed"),
    ("test-passed", "paper"),
)
WINDOWS: tuple[str, ...] = ("dev", "test")
INSIGHT_KINDS: tuple[str, ...] = (
    "observation", "hypothesis", "data-wish", "feature-wish", "risk", "synthesis",
)
# The owner's risk appetite on the dev window, on top of the five P7a D8 conditions.
#
# 0.95 -> 0.90 on 2026-10-07, on the owner's instruction in the session that produced
# LAB_LUCK_GATE_PLAN.md: "this 0.95 threshold is too high man. my risk appetite is 0.90".
# Decision D1 records why this is the lever that moved and the other one is not.
#
# Measured on the committed lab at the time of the change, N = 110 dev trials:
#   at 0.95 nothing is eligible -- the deadlock, 110 trials and 0 promotions;
#   at 0.90 three candidates clear the whole gate once the owner's other change of the same day
#           lands (tuning.MAX_DRAWDOWN 0.15 -> 0.20, Decision D6, phase 8): M0022-W-TV14 (0.912),
#           M0022-W-TV16 (0.916) and M0020-W-NOSTOP (0.913 at N=110, 19.3% drawdown);
#   three near misses stay out, each for its own reason: M0007-N20-RAW clears the new drawdown
#           bar at 19.6% but scores 0.898 re-evaluated at N=110; M0019-RAW20-S25 draws down 20.7%,
#           outside even the new bar; M0001-TV10 does not beat SPY TR.
#
# This is a loosening of a gate that exists to prevent self-deception, and it is the owner's
# call to make. What it is not is a drift: it is one number, cited, dated, measured, and undone
# by editing this line back to 0.95. Decision D1b records the consequence it does not fix --
# at 0.90 the best candidate still sinks below the bar at N ~ 200, about four more Sera nights.
DSR_MIN = 0.90

# The luck label, as it is written into ``trials.failed``. Derived from DSR_MIN rather than
# retyped, so the label and the threshold can never disagree.
#
# ``trials`` is append-only: 110 recorded rows carry the OLD text "DSR >= 0.95". Recognising the
# luck label by ``label == DSR_LABEL`` therefore stops working the moment DSR_MIN moves, and a
# luck-only rejection would read as an owner failure and could never be reconsidered. Every
# reader must use ``is_luck_label`` instead, which matches the label's *shape* and so matches
# every threshold the lab has ever recorded.
LUCK_LABEL_PREFIX = "DSR >= "
DSR_LABEL = f"{LUCK_LABEL_PREFIX}{DSR_MIN:.2f}"


def is_luck_label(label: str) -> bool:
    """True for a luck-test failure label recorded under *any* threshold this lab has used.

    ``DSR_LABEL`` is the one written today; ``"DSR >= 0.95"`` is written on 110 recorded rows.
    Both are the same condition at different thresholds, and the five P7a D8 conditions are
    never of this shape, so the prefix is an exact discriminator.
    """
    return str(label).startswith(LUCK_LABEL_PREFIX)


# The N the luck test deflates by, resolved through ``npolicy.effective_n``. ONE constant: the
# only place in the lab where the multiple-testing count is decided, read at call time by
# ``gate`` below rather than frozen into a row at run time.
#
# **Shipped as "all-trials" deliberately, and NOT inherited from npolicy's own default.**
# Decision D1: the owner moved the threshold, not N, so N stays at every dev trial -- 110 today,
# the same number ``runner.trial_rows`` used before this phase existed. The policy module, its
# correlation evidence (mean pairwise rho 0.595 across the 110 recorded curves, effective N 2.4,
# 23 distinct methods) and phase 5's read-only ``lab luck`` are all built, so the second lever
# is measured and ready -- it is simply not pulled.
#
# Set explicitly, and always passed as an argument to ``npolicy.effective_n``, so that nobody has
# to reason about which module's default wins. ``test_the_shipped_defaults_reproduce_todays_n``
# holds it to 110.
DSR_POLICY = "all-trials"

# The first words of the analysis section `record_promotion` appends, and its idempotence key:
# a method whose analysis already names this roster id has been recorded and is not recorded twice.
PROMOTION_MARKER = "Promoted to the paper roster as "

#: How a roster entry was admitted (lab-luck-gate R4, D3): either the lab's own test window
#: passed the variant, or the owner took it anyway and said why. The same two words as
#: ``paper.roster.BASES``, which is the roster's side of the same bridge.
PROMOTION_BASES: tuple[str, ...] = ("test-passed", "owner-override")


def promotion_basis(status: str) -> str:
    """The admission basis a method's lab status implies at the moment it is promoted.

    ``'test-passed'`` only from the two statuses on the far side of the lab's own test window;
    everything else -- ``'rejected'`` included -- is an ``'owner-override'``. An override is
    allowed (the paper roster's admission rule is not the lab's gate: ``paper.roster``'s module
    docstring, plan Decisions D3) and must state a reason, which is what
    :func:`record_promotion` enforces.
    """
    return "test-passed" if status in ("test-passed", "paper") else "owner-override"


def _quoted(values: Iterable[str]) -> str:
    return ", ".join(f"'{v}'" for v in values)


# The insights table and its two triggers, one statement each. ``_SCHEMA`` creates them on a new
# database; ``_migrate`` recreates them when it rebuilds a v1 table (SQLite cannot alter a CHECK).
_INSIGHTS_TABLE = f"""CREATE TABLE IF NOT EXISTS insights (
    id        INTEGER PRIMARY KEY AUTOINCREMENT,
    kind      TEXT NOT NULL CHECK (kind IN ({_quoted(INSIGHT_KINDS)})),
    title     TEXT NOT NULL CHECK (length(trim(title)) > 0),
    body      TEXT NOT NULL CHECK (length(trim(body)) > 0),
    method_id TEXT REFERENCES methods(id),
    added     TEXT NOT NULL
)"""
_INSIGHTS_TRIGGERS: tuple[str, ...] = (
    """CREATE TRIGGER IF NOT EXISTS insights_no_update BEFORE UPDATE ON insights
BEGIN SELECT RAISE(ABORT, 'insights are append-only: add a newer one instead'); END""",
    """CREATE TRIGGER IF NOT EXISTS insights_no_delete BEFORE DELETE ON insights
BEGIN SELECT RAISE(ABORT, 'insights are append-only'); END""",
)

# The trial_moments side table and its two triggers, one statement each. ``_SCHEMA`` creates them
# on a new database; ``_migrate`` creates them on a v2 database. They are deliberately separate
# constants rather than inline SQL so that there is exactly one definition of the v3 table: the
# migration cannot drift from the fresh schema, and the tests can strip them back out to
# reconstruct a v2 database from ``_SCHEMA`` itself.
#
# ``trial_n`` is the primary key, so a trial has at most one moments row and ``INSERT`` is the
# whole lifecycle: the two triggers below refuse UPDATE and DELETE exactly as ``trials`` does.
# ``var_trials`` is nullable because a lab with fewer than two dev Sharpes has no variance to
# deflate by -- the same condition under which ``trials.dsr`` is NULL.
_MOMENTS_TABLE = """CREATE TABLE IF NOT EXISTS trial_moments (
    trial_n    INTEGER PRIMARY KEY REFERENCES trials(n),
    sr_daily   REAL NOT NULL,
    t          INTEGER NOT NULL CHECK (t >= 2),
    skew       REAL NOT NULL,
    kurt       REAL NOT NULL,
    var_trials REAL,
    n_at_run   INTEGER NOT NULL CHECK (n_at_run >= 1),
    measured   TEXT NOT NULL CHECK (length(trim(measured)) > 0)
)"""
_MOMENTS_TRIGGERS: tuple[str, ...] = (
    """CREATE TRIGGER IF NOT EXISTS trial_moments_no_update BEFORE UPDATE ON trial_moments
BEGIN SELECT RAISE(ABORT, 'trial_moments are append-only: a moments row is never updated'); END""",
    """CREATE TRIGGER IF NOT EXISTS trial_moments_no_delete BEFORE DELETE ON trial_moments
BEGIN SELECT RAISE(ABORT, 'trial_moments are append-only: a moments row is never deleted'); END""",
)

_SCHEMA = f"""
CREATE TABLE IF NOT EXISTS meta (
    key   TEXT PRIMARY KEY,
    value TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS transitions (
    src TEXT NOT NULL,
    dst TEXT NOT NULL,
    PRIMARY KEY (src, dst)
);

CREATE TABLE IF NOT EXISTS methods (
    id            TEXT PRIMARY KEY CHECK (id GLOB 'M[0-9][0-9][0-9][0-9]' OR id GLOB 'H-*'),
    name          TEXT NOT NULL CHECK (length(trim(name)) > 0),
    family        TEXT NOT NULL CHECK (length(trim(family)) > 0),
    parent_id     TEXT REFERENCES methods(id),
    source_kind   TEXT NOT NULL CHECK (source_kind IN ({_quoted(SOURCE_KINDS)})),
    source_ref    TEXT NOT NULL DEFAULT '',
    hypothesis    TEXT NOT NULL CHECK (length(trim(hypothesis)) > 0),
    status        TEXT NOT NULL CHECK (status IN ({_quoted(STATUSES)})),
    analysis      TEXT NOT NULL DEFAULT '',
    verdict       TEXT NOT NULL DEFAULT '',
    blocked_on    TEXT NOT NULL DEFAULT '',
    source_sha    TEXT,
    created       TEXT NOT NULL,
    updated       TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS trials (
    n                 INTEGER PRIMARY KEY AUTOINCREMENT,
    method_id         TEXT NOT NULL REFERENCES methods(id),
    candidate_id      TEXT NOT NULL,
    config_digest     TEXT NOT NULL,
    config_text       TEXT NOT NULL,
    rules_id          TEXT NOT NULL,
    allocator_id      TEXT NOT NULL,
    window            TEXT NOT NULL CHECK (window IN ({_quoted(WINDOWS)})),
    start             TEXT NOT NULL,
    "end"             TEXT NOT NULL,
    store_fingerprint TEXT NOT NULL,
    git_sha           TEXT NOT NULL,
    run_at            TEXT NOT NULL,
    total_return      REAL,
    cagr              REAL,
    max_drawdown      REAL,
    profit_factor     REAL,
    trades            INTEGER NOT NULL,
    sharpe            REAL,
    exposure          REAL,
    turnover          REAL,
    worst_year        INTEGER,
    worst_year_return REAL,
    spy_tr_return     REAL,
    spy_tr_cagr       REAL,
    mar               REAL,
    failed            TEXT NOT NULL,
    eligible          INTEGER NOT NULL CHECK (eligible IN (0, 1)),
    dsr               REAL,
    n_trials_at_run   INTEGER NOT NULL,
    curve_json        TEXT NOT NULL,
    UNIQUE (config_digest, window),
    UNIQUE (candidate_id, window)
);

CREATE TABLE IF NOT EXISTS ideas_seen (
    key       TEXT PRIMARY KEY CHECK (key GLOB '?*:?*'),
    method_id TEXT REFERENCES methods(id),
    note      TEXT NOT NULL DEFAULT '',
    added     TEXT NOT NULL
);

{_INSIGHTS_TABLE};

{_INSIGHTS_TRIGGERS[0]};

{_INSIGHTS_TRIGGERS[1]};

{_MOMENTS_TABLE};

{_MOMENTS_TRIGGERS[0]};

{_MOMENTS_TRIGGERS[1]};

CREATE TRIGGER IF NOT EXISTS trials_no_update BEFORE UPDATE ON trials
BEGIN SELECT RAISE(ABORT, 'trials are append-only: a trial row is never updated'); END;

CREATE TRIGGER IF NOT EXISTS trials_no_delete BEFORE DELETE ON trials
BEGIN SELECT RAISE(ABORT, 'trials are append-only: a trial row is never deleted'); END;

CREATE TRIGGER IF NOT EXISTS methods_no_delete BEFORE DELETE ON methods
BEGIN SELECT RAISE(ABORT, 'methods are never deleted'); END;

CREATE TRIGGER IF NOT EXISTS methods_status_forward BEFORE UPDATE OF status ON methods
WHEN NEW.status <> OLD.status
 AND NOT EXISTS (SELECT 1 FROM transitions WHERE src = OLD.status AND dst = NEW.status)
BEGIN SELECT RAISE(ABORT, 'method status may only move forward along TRANSITIONS'); END;

CREATE TRIGGER IF NOT EXISTS methods_analysis_grows BEFORE UPDATE OF analysis ON methods
WHEN substr(NEW.analysis, 1, length(OLD.analysis)) IS NOT OLD.analysis
BEGIN SELECT RAISE(ABORT, 'analysis is append-only: add to it, never rewrite it'); END;

CREATE TRIGGER IF NOT EXISTS methods_hypothesis_frozen BEFORE UPDATE OF hypothesis ON methods
WHEN OLD.status <> 'idea' AND NEW.hypothesis IS NOT OLD.hypothesis
BEGIN SELECT RAISE(ABORT, 'the hypothesis is frozen once a method leaves idea'); END;

CREATE TRIGGER IF NOT EXISTS methods_source_sha_once BEFORE UPDATE OF source_sha ON methods
WHEN OLD.source_sha IS NOT NULL AND NEW.source_sha IS NOT OLD.source_sha
BEGIN SELECT RAISE(ABORT, 'source_sha is set once, when the method runs'); END;

CREATE TRIGGER IF NOT EXISTS seen_no_delete BEFORE DELETE ON ideas_seen
BEGIN SELECT RAISE(ABORT, 'ideas_seen is append-only'); END;
"""


class LabError(RuntimeError):
    """A lab operation the database's rules refuse (exit 2 in the CLI)."""


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def connect(path: Path = DB_PATH) -> sqlite3.Connection:
    """Open (creating when missing) the lab database with the schema, triggers and transitions,
    migrating an older database to ``SCHEMA_VERSION`` first."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path, timeout=120)  # parallel explorer sessions share one file
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("PRAGMA foreign_keys = ON")
        conn.execute("PRAGMA journal_mode = DELETE")
        with conn:
            conn.executescript(_SCHEMA)
            conn.executemany("INSERT OR IGNORE INTO transitions (src, dst) VALUES (?, ?)", TRANSITIONS)
            conn.execute(
                "INSERT OR IGNORE INTO meta (key, value) VALUES ('schema_version', ?)", (SCHEMA_VERSION,)
            )
        _migrate(conn)
    except BaseException:
        conn.close()
        raise
    return conn


def connect_readonly(path: Path = COMMITTED_DB) -> sqlite3.Connection:
    """Open an existing lab database read-only, exactly as it is: no schema, no migration, no
    write (``mode=ro``). For the snapshot sync guard, which must not touch the committed file."""
    path = Path(path).resolve()
    if not path.is_file():
        raise LabError(f"no lab database at {path}")
    conn = sqlite3.connect(f"{path.as_uri()}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def schema_version(conn: sqlite3.Connection) -> str | None:
    row = conn.execute("SELECT value FROM meta WHERE key = 'schema_version'").fetchone()
    return None if row is None else str(row[0])


def _v1_to_v2(conn: sqlite3.Connection) -> None:
    """Add the ``synthesis`` insight kind. SQLite cannot alter a CHECK, so ``insights`` is
    rebuilt: the v1 table is renamed aside, the v2 table is created from ``_INSIGHTS_TABLE``,
    every row is copied with its id, the AUTOINCREMENT counter is carried over, the v1 table is
    dropped (which drops its triggers and fires none), and the two append-only triggers are
    created on the new table. The triggers come last: while the v1 table exists its triggers hold
    their names, and ``CREATE TRIGGER IF NOT EXISTS`` would silently skip them."""
    seq = conn.execute("SELECT seq FROM sqlite_sequence WHERE name = 'insights'").fetchone()
    conn.execute("ALTER TABLE insights RENAME TO insights_v1")
    conn.execute(_INSIGHTS_TABLE)
    conn.execute(
        "INSERT INTO insights (id, kind, title, body, method_id, added) "
        "SELECT id, kind, title, body, method_id, added FROM insights_v1 ORDER BY id"
    )
    conn.execute("DROP TABLE insights_v1")
    if seq is not None:
        cur = conn.execute(
            "UPDATE sqlite_sequence SET seq = max(seq, ?) WHERE name = 'insights'", (seq[0],)
        )
        if cur.rowcount == 0:
            conn.execute("INSERT INTO sqlite_sequence (name, seq) VALUES ('insights', ?)", (seq[0],))
    for trigger in _INSIGHTS_TRIGGERS:
        conn.execute(trigger)


def _v2_to_v3(conn: sqlite3.Connection) -> None:
    """Add the ``trial_moments`` side table and its two append-only triggers.

    Purely additive, and deliberately so: no ``trials`` row is read, written, rebuilt or
    re-keyed, no column is altered, and no verdict moves. A v2 database that migrates and is then
    never written again differs from its v2 self only by an empty table, two triggers and the
    ``schema_version`` string. (``connect`` has already run ``_SCHEMA`` by the time this is
    called, so in practice these are no-ops; they are issued anyway so the migration is complete
    on its own terms and does not depend on the order ``connect`` happens to use.)
    """
    conn.execute(_MOMENTS_TABLE)
    for trigger in _MOMENTS_TRIGGERS:
        conn.execute(trigger)


def _migrate(conn: sqlite3.Connection) -> None:
    """Bring an older database up to ``SCHEMA_VERSION`` in one transaction under the write lock.

    A ladder: each step moves the database up exactly one version, so a v1 database reaches v3 in
    one open by running both steps in order. v1 -> v2 adds the ``synthesis`` insight kind
    (``_v1_to_v2``); v2 -> v3 adds the ``trial_moments`` side table (``_v2_to_v3``). A version
    this code does not know is refused rather than guessed at.

    Parallel sessions may connect at once, so the version is read again after
    ``BEGIN IMMEDIATE``: only the first one migrates, the others find the current version and do
    nothing.
    """
    if schema_version(conn) == SCHEMA_VERSION:
        return
    begin_immediate(conn)
    try:
        found = schema_version(conn)
        version = found
        if version == "1":
            _v1_to_v2(conn)
            version = "2"
        if version == "2":
            _v2_to_v3(conn)
            version = "3"
        if version != SCHEMA_VERSION:
            raise LabError(
                f"lab database schema version {found!r} is unknown to this code "
                f"(expects {SCHEMA_VERSION})"
            )
        conn.execute("UPDATE meta SET value = ? WHERE key = 'schema_version'", (SCHEMA_VERSION,))
        conn.commit()
    except BaseException:
        conn.rollback()
        raise


# --------------------------------------------------------------------------- methods


def begin_immediate(conn: sqlite3.Connection) -> None:
    """Take the write lock now, so a read-then-write (next id, lab-wide N) is atomic across sessions."""
    if conn.in_transaction:
        conn.commit()
    conn.execute("BEGIN IMMEDIATE")


def get_method(conn: sqlite3.Connection, method_id: str) -> sqlite3.Row | None:
    return conn.execute("SELECT * FROM methods WHERE id = ?", (method_id,)).fetchone()


def next_method_id(conn: sqlite3.Connection) -> str:
    """The next free ``M0001``-style id (one above the highest used)."""
    row = conn.execute(
        "SELECT max(CAST(substr(id, 2) AS INTEGER)) FROM methods WHERE id GLOB 'M[0-9][0-9][0-9][0-9]'"
    ).fetchone()
    return f"M{(row[0] or 0) + 1:04d}"


def add_method(
    conn: sqlite3.Connection,
    *,
    id: str,
    name: str,
    family: str,
    source_kind: str,
    hypothesis: str,
    source_ref: str = "",
    parent_id: str | None = None,
    status: str = "idea",
    analysis: str = "",
    verdict: str = "",
    allow_any_status: bool = False,
) -> None:
    """Insert a method. New lab methods start as ``idea`` or ``registered``; only the seed
    import (``allow_any_status``) inserts a closed historical record directly."""
    if not allow_any_status and status not in ("idea", "registered"):
        raise LabError(f"a new method starts as idea or registered, not {status}")
    if get_method(conn, id) is not None:
        raise LabError(f"method {id} already exists")
    stamp = now_iso()
    conn.execute(
        "INSERT INTO methods (id, name, family, parent_id, source_kind, source_ref, hypothesis, "
        "status, analysis, verdict, created, updated) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (id, name, family, parent_id, source_kind, source_ref, hypothesis, status, analysis, verdict, stamp, stamp),
    )


def update_method(conn: sqlite3.Connection, method_id: str, **values: Any) -> None:
    """Set columns of one method (the triggers still apply) and bump ``updated``."""
    allowed = {"status", "verdict", "blocked_on", "source_sha", "hypothesis", "name", "family",
               "source_kind", "source_ref", "parent_id"}
    bad = set(values) - allowed
    if bad:
        raise LabError(f"cannot set {sorted(bad)} through update_method")
    if get_method(conn, method_id) is None:
        raise LabError(f"no method {method_id}")
    cols = ", ".join(f"{k} = ?" for k in values)
    try:
        conn.execute(
            f"UPDATE methods SET {cols}, updated = ? WHERE id = ?",
            (*values.values(), now_iso(), method_id),
        )
    except sqlite3.IntegrityError as e:
        raise LabError(f"{method_id}: {e}") from e


def append_analysis(conn: sqlite3.Connection, method_id: str, text: str) -> None:
    """Append a dated section to a method's analysis (never rewrites what is there).

    The author's own ``### <today>`` opening line counts as the section head. A note written
    that way used to land under a second, identical heading, and the site then showed the date
    twice; the analysis is append-only, so the duplicate could never be taken back out.
    """
    row = get_method(conn, method_id)
    if row is None:
        raise LabError(f"no method {method_id}")
    text = text.strip()
    if not text:
        raise LabError("empty analysis")
    head = f"### {date.today().isoformat()}"
    if text != head and not text.startswith(head + "\n"):
        text = head + "\n\n" + text
    new = (row["analysis"] + "\n\n" if row["analysis"] else "") + text + "\n"
    conn.execute(
        "UPDATE methods SET analysis = ?, updated = ? WHERE id = ?", (new, now_iso(), method_id)
    )


def record_promotion(
    conn: sqlite3.Connection,
    *,
    method_id: str,
    strategy_id: str,
    candidate_id: str,
    object_name: str,
    spec_digest: str,
    retired_id: str | None = None,
    move_status: bool = True,
    basis: str | None = None,
    reason: str = "",
) -> str:
    """Record that ``candidate_id`` became paper roster entry ``strategy_id``; return the status.

    The lab is append-only (§1), so a promotion is *added*, never stamped over anything:

    - ``analysis`` grows by one dated section (``methods_analysis_grows`` permits only growth);
    - one ``insights`` row is appended (the journal table refuses UPDATE and DELETE outright);
    - ``status`` moves to ``paper`` **only** along the edge ``TRANSITIONS`` already has,
      ``('test-passed', 'paper')``. A method already at ``paper`` is left alone. From any other
      status this raises LabError, because there is no edge and inventing one would make the
      lab's own vocabulary mean less. Pass ``move_status=False`` to record the promotion and
      leave the status where it is -- the honest shape for a roster that is taking a method the
      lab has not passed (the roster's admission rule is not the lab's gate: plan Decisions D3).

    **The basis (lab-luck-gate R4, D3).** ``basis`` is how the roster admitted the variant:
    ``'test-passed'`` when the lab's own test window passed it, ``'owner-override'`` when the
    owner took it anyway. Left ``None`` it is derived from the method's current status by
    :func:`promotion_basis`, which is right for every caller that has just read that status. An
    ``'owner-override'`` **must** carry a one-line ``reason``; without one this raises, because an
    unexplained override is exactly the silent divergence this argument exists to end. The basis,
    the reason and the status at admission are written into the analysis section, so the lab
    records the same fact ``paper.roster.LAB_PROVENANCE`` states on the roster's side.

    ``hypothesis``, ``verdict``, ``parent_id`` and above all ``source_sha`` are never written.
    No ``trials`` row is inserted: a promotion is not a backtest and must not move the lab's N.

    Idempotent: a method whose ``analysis`` already carries ``PROMOTION_MARKER`` followed by
    ``strategy_id`` is already recorded, and this writes nothing and returns the current status --
    including when this call's ``basis`` or ``reason`` differ, because what was recorded is what
    was true on the night of the admission. That is what lets the command be re-run to repair a
    half-finished promotion, since the roster (Neon) and the lab (SQLite) cannot share one
    transaction.

    The caller holds the transaction (``begin_immediate`` / ``with conn``), as every other writer
    in this module does.
    """
    row = get_method(conn, method_id)
    if row is None:
        raise LabError(f"no method {method_id}")
    if not candidate_id.startswith(method_id):
        raise LabError(f"candidate {candidate_id!r} does not belong to method {method_id}")
    status = str(row["status"])
    if PROMOTION_MARKER + f"`{strategy_id}`" in row["analysis"]:
        return status

    if move_status and status != "paper":
        if (status, "paper") not in TRANSITIONS:
            raise LabError(
                f"{method_id} is {status!r} and the lab's TRANSITIONS have no edge "
                f"{status!r} -> 'paper'; only 'test-passed' reaches 'paper'. The roster may still "
                f"take this method -- its admission rule is not the lab's gate -- but say so: "
                f"re-run with --lab-status-stays, which records the promotion and leaves the "
                f"status alone."
            )

    basis = promotion_basis(status) if basis is None else basis
    if basis not in PROMOTION_BASES:
        raise LabError(f"basis {basis!r} is not one of {PROMOTION_BASES}")
    reason = reason.strip()
    if basis == "owner-override" and not reason:
        raise LabError(
            f"{method_id} is {status!r}, so the roster is admitting a method the lab has not "
            f"passed. That is allowed -- the roster's admission rule is not the lab's gate -- but "
            f"the basis goes on the record: pass reason='the one line that says why'."
        )

    retired = "" if retired_id is None else f", replacing `{retired_id}` (retired the same moment)"
    why = f" {reason}" if reason else ""
    body = (
        f"{PROMOTION_MARKER}`{strategy_id}`{retired}.\n\n"
        f"Variant: `{candidate_id}`. Roster object: `{object_name}`. "
        f"Frozen spec digest: `{spec_digest}`.\n\n"
        f"Lab status at admission: `{status}`. Basis: `{basis}`.{why}\n\n"
        f"The roster row is `status='active'` with no `paper_start`: the next paper night freezes "
        f"the spec and starts its own clock, so the paper record begins at the promotion and "
        f"claims nothing earlier. `backtest.registry.REGISTRY` was not appended to -- a promoted "
        f"method reaches the roster through the roster's own resolver, so the lab's "
        f"multiple-testing count is unchanged by this."
    )
    append_analysis(conn, method_id, "# Promotion\n\n" + body)
    # The journal note is read by the owner, not an auditor: plain words, no digests or column names.
    # The technical record above stays in the method's analysis.
    replacing = "" if retired_id is None else f", taking the place of {retired_id}"
    add_insight(
        conn,
        kind="observation",
        title=f"{row['name']} starts paper trading as {strategy_id}",
        body=(
            f"Sera picked this method to trade with pretend money every night, under the short "
            f"name {strategy_id}{replacing}. Its rules are now locked, and its record starts from "
            f"its first night on paper, so nothing before that counts. Month by month against SPY "
            f"is how it earns trust."
        ),
        method_id=method_id,
    )
    if move_status and status != "paper":
        update_method(conn, method_id, status="paper")
        return "paper"
    return status



# --------------------------------------------------------------------------- trials


@dataclass(frozen=True)
class TrialRow:
    """One ``trials`` row before insertion (``n`` is assigned by SQLite)."""

    method_id: str
    candidate_id: str
    config_digest: str
    config_text: str
    rules_id: str
    allocator_id: str
    window: str
    start: str
    end: str
    store_fingerprint: str
    git_sha: str
    run_at: str
    total_return: float | None
    cagr: float | None
    max_drawdown: float | None
    profit_factor: float | None
    trades: int
    sharpe: float | None
    exposure: float | None
    turnover: float | None
    worst_year: int | None
    worst_year_return: float | None
    spy_tr_return: float | None
    spy_tr_cagr: float | None
    mar: float | None
    failed: str  # "; "-joined failure labels, "" when eligible
    eligible: bool
    dsr: float | None
    n_trials_at_run: int
    curve_json: str


TRIAL_COLUMNS: tuple[str, ...] = tuple(f.name for f in fields(TrialRow))


def has_trial(conn: sqlite3.Connection, config_digest: str, window: str) -> bool:
    return (
        conn.execute(
            "SELECT 1 FROM trials WHERE config_digest = ? AND window = ?", (config_digest, window)
        ).fetchone()
        is not None
    )


def insert_trials(conn: sqlite3.Connection, rows: Sequence[TrialRow]) -> list[int]:
    """Insert trials (the caller holds the transaction). Returns their trial numbers."""
    out: list[int] = []
    cols = ", ".join(f'"{c}"' for c in TRIAL_COLUMNS)
    marks = ", ".join("?" for _ in TRIAL_COLUMNS)
    for r in rows:
        if r.window not in WINDOWS:
            raise LabError(f"unknown window {r.window!r}")
        if has_trial(conn, r.config_digest, r.window):
            raise LabError(
                f"{r.candidate_id}: this configuration already has a {r.window} trial "
                "(no re-rolls: a change of params, rules or allocator is a new method)"
            )
        values = [getattr(r, c) for c in TRIAL_COLUMNS]
        values[TRIAL_COLUMNS.index("eligible")] = int(r.eligible)
        cur = conn.execute(f"INSERT INTO trials ({cols}) VALUES ({marks})", values)
        out.append(int(cur.lastrowid))
    return out


def dev_trial_count(conn: sqlite3.Connection) -> int:
    return int(conn.execute("SELECT count(*) FROM trials WHERE window = 'dev'").fetchone()[0])


def test_looks(conn: sqlite3.Connection) -> int:
    return int(conn.execute("SELECT count(*) FROM trials WHERE window = 'test'").fetchone()[0])


def dev_daily_sharpes(conn: sqlite3.Connection) -> list[float]:
    """Every dev trial's daily Sharpe (annualized Sharpe / sqrt(252)), for the DSR's variance."""
    rows = conn.execute("SELECT sharpe FROM trials WHERE window = 'dev' AND sharpe IS NOT NULL")
    return [float(r[0]) / math.sqrt(252) for r in rows]


def trials_of(conn: sqlite3.Connection, method_id: str) -> list[sqlite3.Row]:
    return conn.execute("SELECT * FROM trials WHERE method_id = ? ORDER BY n", (method_id,)).fetchall()


def best_dev_eligible(
    conn: sqlite3.Connection,
    method_id: str,
    *,
    at: Gate | None = None,
    policy: str | None = None,
) -> sqlite3.Row | None:
    """The method's best eligible dev trial by MAR -- the one variant design §3 pre-registers.

    Highest MAR wins and a tie breaks on the trial number, so the answer is exactly one row and
    the same row every time: "one per method" is a property of this query, not of the caller.

    Only ``window = 'dev'`` is considered. A test trial is the out-of-sample check on a
    configuration this query already chose, so letting one back in here would let a test number
    decide what gets tested. A trial with no MAR is never the answer either -- an eligible trial
    always has one, because ``beats SPY TR`` is among the conditions it passed, so a NULL here
    means a row that cannot be compared rather than a row that compares badly.

    **Eligible means ``verdict(...).eligible``, not the frozen ``eligible`` column.** The verdict
    is read under one threshold and one N at evaluation time (plan Decisions D1/D2), so two
    trials recorded six weeks apart are ranked against the same bar instead of against whichever
    bar happened to be in force on each run date. Nothing else about the choice changed: the
    ordering is still MAR then ``n``. Selection is still **never** on DSR -- a higher-DSR variant
    does not outrank a higher-MAR one, which is why ``M0022-W-TV14`` (MAR 0.857, DSR 0.912) is
    the answer for M0022 and ``M0022-W-TV16`` (MAR 0.816, DSR 0.916) is not.

    The gate is resolved once, after the rows are fetched, so a method with no comparable dev
    trial costs no estimator read at all.

    ``at`` passes in an already-resolved gate. A caller looping over many methods -- ``lab
    status`` walks all 23 with dev trials -- should resolve ``gate(conn)`` once and pass it to
    every call, so the estimator provably runs once rather than relying on ``_GATE_CACHE``.
    ``policy`` is for phase 5's read-only comparison; the promotion path passes neither.

    None when the method has no eligible dev trial at all.
    """
    rows = conn.execute(
        "SELECT * FROM trials WHERE method_id = ? AND window = 'dev' AND mar IS NOT NULL "
        "ORDER BY mar DESC, n ASC",
        (method_id,),
    ).fetchall()
    if not rows:
        return None
    g = gate(conn, policy) if at is None else at
    for row in rows:
        if verdict(conn, row, at=g).eligible:
            return row
    return None


# --------------------------------------------------------------------------- the derived verdict


@dataclass(frozen=True)
class Gate:
    """The luck bar in force right now: the policy name and the N it resolved to.

    The threshold is not in here: it is ``DSR_MIN``, one module constant, and there is no policy
    over it.
    """

    n: int
    policy: str


@dataclass(frozen=True)
class Verdict:
    """One dev trial's eligibility as it reads *now*.

    ``failed`` is the ordered tuple of failure labels (``trials.failed`` is the same list,
    "; "-joined), owner conditions first in ``dev.FAILURE_LABELS`` order and the luck label last.
    ``eligible`` is ``failed == ()``, always.

    ``derived`` is True when the DSR in this verdict was **evaluated at ``n``**, and False when
    it could not be -- in which case ``dsr`` is None and the luck label is among ``failed``. It
    never means "the recorded columns were returned verbatim": nothing is returned verbatim.
    """

    dsr: float | None
    failed: tuple[str, ...]
    eligible: bool
    n: int
    policy: str
    derived: bool


# Resolved gates, keyed by (database file, policy, the dev trial set's fingerprint).
#
# ``npolicy.effective_n`` decodes 110 month-end curves out of ``curve_json`` and takes an
# eigendecomposition of their correlation matrix for the participation ratio. ``lab status``
# calls ``best_dev_eligible`` once per method with dev trials -- 23 of them -- so without this
# the estimator would run 23 times for one command.
#
# The key is derived from the *content*, not from the connection: ``sqlite3.Connection`` supports
# neither weak references nor attributes, so there is nowhere on it to hang a cache and nothing
# safe to key on (``id()`` is reused after a connection is freed). ``trials`` is append-only --
# the ``trials_no_update`` and ``trials_no_delete`` triggers see to that -- so for one database
# file the triple (dev row count, highest dev trial number, distinct dev methods) pins the dev
# trial set exactly, and any insert changes it. That makes the cache self-invalidating.
#
# In-memory databases are never cached: their path is "" and two of them could otherwise collide
# on identical counts.
_GATE_CACHE: dict[tuple[str, str, int, int, int], Gate] = {}
_GATE_CACHE_MAX = 32


def _gate_key(conn: sqlite3.Connection, policy: str) -> tuple[str, str, int, int, int] | None:
    """A content-derived cache key, or None when this database must not be cached."""
    path = ""
    for _seq, name, file in conn.execute("PRAGMA database_list"):
        if name == "main":
            path = str(file or "")
            break
    if not path:
        return None  # ":memory:" and temporary databases
    count, high, methods = conn.execute(
        "SELECT count(*), coalesce(max(n), 0), count(DISTINCT method_id) "
        "FROM trials WHERE window = 'dev'"
    ).fetchone()
    return (path, policy, int(count), int(high), int(methods))


def gate(conn: sqlite3.Connection, policy: str | None = None) -> Gate:
    """The N the luck test deflates by on this database, under ``policy`` (default ``DSR_POLICY``).

    Memoised per (database file, policy, dev trial set) -- see ``_GATE_CACHE``. A caller judging
    many trials should still resolve it once and pass it down as ``verdict(..., at=g)`` or
    ``best_dev_eligible(..., at=g)``: that is explicit and does not depend on the cache being
    warm or on the key being right.

    The policy is always passed to ``npolicy.effective_n`` explicitly; this module never relies
    on that module's own default.

    ``npolicy`` is imported here rather than at module scope, the way ``snapshot`` imports its
    dependencies: importing ``store`` is something the whole lab does, and it should never pull
    in the estimator's curve arithmetic, nor create a cycle if ``npolicy`` ever needs ``store``.
    """
    from seer_engine.lab import npolicy

    name = DSR_POLICY if policy is None else policy
    key = _gate_key(conn, name)
    if key is not None:
        hit = _GATE_CACHE.get(key)
        if hit is not None:
            return hit
    count = npolicy.effective_n(conn, name)
    resolved = Gate(n=int(count.n), policy=str(count.policy))
    if key is not None:
        if len(_GATE_CACHE) >= _GATE_CACHE_MAX:
            _GATE_CACHE.clear()  # a long-lived process never accumulates stale paths
        _GATE_CACHE[key] = resolved
    return resolved


def pending_gate(
    conn: sqlite3.Connection, method_id: str, pending: int, *, policy: str | None = None
) -> Gate:
    """The gate a batch of ``pending`` new dev trials for ``method_id`` will be judged under.

    ``gate`` counts what the database holds, and the batch is not in it yet: the DSR and the
    eligibility have to be decided *before* ``insert_trials`` runs, because ``trials`` is
    append-only and a row is written exactly once. So the count is projected forward over the
    batch, in the unit the policy counts in:

    - ``all-trials``  N + ``pending``  -- every new row is another trial. Under the shipped
      policy this is exactly ``dev_trial_count(conn) + len(results)``, the expression
      ``runner.trial_rows`` used before this phase, so a new ``lab run`` is deflated by the
      number it has always been deflated by.
    - ``methods``     N + 1 when ``method_id`` has no dev trial yet, N otherwise -- a batch of
      variants of one method is one method. The participation-ratio floor is not re-measured
      against curves that do not exist yet; it is a floor, and this projection only ever sits
      on or above it.
    - ``effective``   N -- the measured independence of curves that have not been recorded
      cannot be projected, so the batch is judged under the independence already measured.
    - anything else   N + ``pending``, the most punishing of the three. A policy name this
      module does not recognise must not quietly deflate a new trial by less than the row count.

    Every branch returns at least ``gate(conn).n``, so a trial recorded today is never deflated
    by a smaller N than one recorded yesterday under the same policy.
    """
    g = gate(conn, policy)
    if pending <= 0:
        return g
    if g.policy == "effective":
        return g
    if g.policy == "methods":
        seen_already = conn.execute(
            "SELECT 1 FROM trials WHERE method_id = ? AND window = 'dev' LIMIT 1", (method_id,)
        ).fetchone()
        return g if seen_already is not None else Gate(n=g.n + 1, policy=g.policy)
    return Gate(n=g.n + pending, policy=g.policy)


# The one D8 condition with no number in it, and therefore the one recorded label that can never
# go stale: it asks whether a human hand-picked a parameter of the candidate, not whether a metric
# cleared a bar. ``owner_failures`` carries it from the recorded string instead of re-deriving it,
# because there is no column to re-derive it from.
#
# A literal rather than ``dev.FAILURE_LABELS[-1]``: ``store`` deliberately imports ``dev`` only
# inside the functions that need it, and a module-level constant would pull the whole backtest
# package into every ``import store``.
# ``test_the_owner_inputs_label_is_the_one_dev_still_writes`` pins the two together.
OWNER_INPUTS_LABEL = "owner inputs"


def recorded_labels(failed: str) -> tuple[str, ...]:
    """A recorded ``trials.failed`` string, split into its labels. **History, not a verdict.**

    Every label it returns names the threshold in force on the trial's **run date**: 110 recorded
    rows say ``"DSR >= 0.95"`` and ``"max DD <= 15%"``, and the live bars are 0.90 and 20%. Use
    this to display what the lab said at the time, or to ask whether a particular condition was
    recorded; never to decide what a trial is today. ``owner_failures`` is that.
    """
    return tuple(f for f in str(failed or "").split("; ") if f)


def owner_failures(trial: Mapping[str, Any] | sqlite3.Row) -> tuple[str, ...]:
    """The five P7a D8 conditions ``trial`` misses **as they read now**, in FAILURE_LABELS order.

    **Four re-derived, one carried.** Four of the five are thresholds over numbers ``trials``
    already records, so they are recomputed here from the recorded columns against the live
    constants -- the same comparisons ``dev.py`` makes, on the same values:

    - ``beats SPY TR``  ``total_return > spy_tr_return``
    - ``max DD``        ``max_drawdown <= tuning.MAX_DRAWDOWN`` (0.20 since 2026-10-07, phase 8)
    - ``PF``            ``profit_factor >= tuning.MIN_PROFIT_FACTOR``
    - ``trades``        ``trades >= dev._MIN_TRADES``

    The fifth, ``owner inputs``, is **carried from the recorded ``failed`` string**, because it is
    the one condition that is not a threshold: it asks whether a human hand-picked a parameter of
    the *candidate* (``dev.candidate_owner_inputs``), there is no column to recompute it from, and
    no constant re-decides it -- so its recorded label can never go stale.

    **Why nothing is parsed out of ``failed`` for the other four.** ``trials`` is append-only, so
    a recorded label names the bar in force on its run date. All 110 recorded rows carry
    ``"max DD <= 15%"``, and the owner moved that bar to 20% on 2026-10-07; reading the condition
    back out of the string would freeze every recorded trial at the bar it was judged by and
    ``M0020-W-NOSTOP`` (19.3%) could never become eligible. That is precisely the "verdicts mix
    bars" defect R2 names, and it is why this function takes the row rather than the string.

    The returned labels are the **live** ones (``dev.FAILURE_LABELS``), so a caller printing them
    states today's bar, not the one the row was judged by.
    """
    from seer_engine.backtest import dev, tuning

    spy, drawdown, pf, trades, owner = dev.FAILURE_LABELS
    total, bench = trial["total_return"], trial["spy_tr_return"]
    dd, factor, count = trial["max_drawdown"], trial["profit_factor"], trial["trades"]
    out: list[str] = []
    if total is None or bench is None or float(total) <= float(bench):
        out.append(spy)
    if dd is None or float(dd) > tuning.MAX_DRAWDOWN:
        out.append(drawdown)
    if factor is None or float(factor) < tuning.MIN_PROFIT_FACTOR:
        out.append(pf)
    if count is None or int(count) < dev._MIN_TRADES:
        out.append(trades)
    if owner in recorded_labels(trial["failed"]):
        out.append(owner)
    return tuple(out)


def sr_star(n_trials: int, var_trials: float) -> float:
    """The deflated Sharpe's daily hurdle ``SR*`` at ``n_trials`` independent looks.

    ``dev.deflated_sharpe``'s own ``sr_star`` line, with its own ``_EULER_GAMMA``, isolated so
    the hurdle can be asked for at an N no trial was ever run at. Nothing is re-derived: this is
    the formula being inverted, not a second opinion about it.
    """
    from statistics import NormalDist

    from seer_engine.backtest import dev

    normal = NormalDist()
    expected_max = (1 - dev._EULER_GAMMA) * normal.inv_cdf(1 - 1 / n_trials) + (
        dev._EULER_GAMMA * normal.inv_cdf(1 - 1 / (n_trials * math.e))
    )
    return math.sqrt(var_trials) * expected_max


def recover_dsr(
    *,
    sharpe_daily: float,
    dsr_at_run: float,
    n_at_run: int,
    var_trials: float,
    n_trials: int,
) -> float | None:
    """A recorded DSR re-evaluated at ``n_trials`` looks. Pure; reads and writes nothing.

    ``DSR = Phi((SR - SR*(N)) * k)`` where ``k = sqrt(t-1)/sqrt(radicand)`` does **not** depend on
    N. ``t``, the skew and the kurtosis are not in ``trials`` -- that is what ``trial_moments``
    exists to fix going forward -- so ``k`` is recovered by inverting a known DSR at the N it
    belongs to, and the answer is the same ``Phi`` with a different ``SR*``.

    **At ``n_trials == n_at_run`` this returns ``dsr_at_run`` exactly, for any ``var_trials``** --
    the two ``SR*`` terms are the same term and ``k`` cancels. That identity is what makes this a
    *re-reading* of the record rather than a second estimate of it.

    **Monotone in N** for any ``dsr_at_run`` above a half: ``SR*`` rises with N and ``k > 0``
    there, so the DSR falls as the search widens. That is R1's ratchet, as arithmetic.

    None where the inversion is undefined: fewer than two looks at either N, a non-positive
    ``var_trials``, a non-finite input, a DSR of exactly 0 or 1 (``Phi^-1`` has no value there),
    or an SR sitting exactly on the hurdle at ``n_at_run`` (``k`` would divide by zero -- the
    record says the trial was exactly at the bar and says nothing about its ``k``).
    """
    from statistics import NormalDist

    if n_trials < 2 or n_at_run < 2 or var_trials <= 0:
        return None
    if not all(math.isfinite(x) for x in (sharpe_daily, dsr_at_run, var_trials)):
        return None
    if not 0.0 < dsr_at_run < 1.0:
        return None
    gap = sharpe_daily - sr_star(n_at_run, var_trials)
    if gap == 0.0:
        return None
    normal = NormalDist()
    k = normal.inv_cdf(dsr_at_run) / gap
    return normal.cdf((sharpe_daily - sr_star(n_trials, var_trials)) * k)


def dev_sharpe_variance(conn: sqlite3.Connection) -> float | None:
    """The variance of the dev trials' daily Sharpes **as the lab stands now**. None below two.

    The same expression ``runner.trial_rows`` deflates by, so the dispersion the expected maximum
    is drawn from is the dispersion of the search as it actually is -- which is the point: the
    gate describes today on both axes, the count of looks and their spread.
    """
    import statistics

    sharpes = dev_daily_sharpes(conn)
    return statistics.variance(sharpes) if len(sharpes) >= 2 else None


def dsr_at(
    conn: sqlite3.Connection, trial: Mapping[str, Any] | sqlite3.Row, n_trials: int
) -> float | None:
    """``trial``'s deflated Sharpe **at ``n_trials`` looks**, or None when it cannot be had.

    Two routes to one number, in order of exactness:

    1. **From ``trial_moments``**, when ``lab run`` recorded them (phase 2) or ``lab remeasure``
       recovered them (phase 3). The same ``dev.deflated_sharpe`` on the trial's own measured
       ``sr_daily``, ``t``, ``skew`` and ``kurt``. An exact recomputation.

    2. **By inverting the recorded ``dsr``** at its own ``n_trials_at_run`` and re-evaluating at
       ``n_trials`` (``recover_dsr``). Exact arithmetic on recorded data -- no backtest is re-run
       and nothing is estimated -- and it returns the recorded number unchanged when the trial
       was already judged at ``n_trials``.

    **Both routes use ``dev_sharpe_variance(conn)`` -- the trial-Sharpe variance as the lab stands
    now -- and never the ``var_trials`` recorded beside the trial.** The function body reads that
    variance **once, before either route branches**, precisely so the two cannot drift apart
    again; ``moments["var_trials"]`` is not referenced anywhere in this function. That is
    deliberate (**Decision D12**) and it is the same principle as using the gate's N rather than
    ``n_trials_at_run``: the deflated Sharpe asks "how extreme is this Sharpe against the maximum
    of N draws from the trial-Sharpe distribution", and **both** N and that distribution describe
    the search as it is today. Pairing today's N with a variance frozen at the run date would mix
    bars on the other axis -- exactly the defect R2 names, one column over. The recorded
    ``var_trials`` stays in ``trial_moments`` as history: it is what the trial *was* judged by,
    and phase 3's reproduction check is what it is for.

    **This is not a stylistic preference; it decides a candidate.** On the 54 P7a seed rows that
    phase 9 re-measures, the recorded ``var_trials`` (2.006691e-04, the P7a search's own) and
    today's (2.395048e-04, all 110 dev trials) differ by enough to move
    ``F9-SPY200M70-MOM30`` from **0.903053** (which clears 0.90) to **0.856651** (which does not).
    Under this function it reads **0.8567** and stays ineligible -- on a luck test it finally
    received rather than on a missing column, which is the whole of R6.

    Measured on the committed lab, the two routes agree to **1.3e-5** across all 56 trials with a
    recorded DSR, which is what makes route 2 a re-reading of the record rather than a second
    opinion about it. A trial recorded *today* reads back exactly as it was recorded, because the
    gate's N and today's variance are then the very values it was judged by.

    **None when the recorded ``dsr`` is NULL**, which is the 54 P7a seed rows by construction
    (``seed.py``: "P7a reported it for one row only"). A luck test that cannot be evaluated is
    a luck test that was not passed -- ``verdict`` appends the luck label, exactly as
    ``runner.trial_rows`` does for a new trial whose DSR comes back None. Nothing is admitted for
    being unmeasurable.
    """
    # ONE variance, read once, used by BOTH routes. `moments["var_trials"]` is deliberately not
    # read anywhere in this function: see Decision D12 and the docstring above.
    var = dev_sharpe_variance(conn)
    if var is None:
        return None
    moments = moments_of(conn, int(trial["n"]))
    if moments is not None:
        from seer_engine.backtest import dev

        return dev.deflated_sharpe(
            float(moments["sr_daily"]),
            n_trials,
            var,
            int(moments["t"]),
            float(moments["skew"]),
            float(moments["kurt"]),
        )
    if trial["dsr"] is None or trial["sharpe"] is None:
        return None
    from seer_engine.backtest.book_runner import TRADING_DAYS

    return recover_dsr(
        sharpe_daily=float(trial["sharpe"]) / math.sqrt(TRADING_DAYS),
        dsr_at_run=float(trial["dsr"]),
        n_at_run=int(trial["n_trials_at_run"]),
        var_trials=var,
        n_trials=n_trials,
    )


def verdict(
    conn: sqlite3.Connection,
    trial: Mapping[str, Any] | sqlite3.Row,
    *,
    at: Gate | None = None,
    policy: str | None = None,
) -> Verdict:
    """``trial``'s eligibility as it reads now: every condition, at the bars in force now.

    ``trial`` is a ``trials`` row (anything addressable by column name, including a
    ``sqlite3.Row``) carrying ``n``, ``dsr``, ``sharpe``, ``failed``, ``n_trials_at_run`` and the
    four metric columns. Meaningful for ``window = 'dev'`` rows: a test trial's DSR is recorded
    and is not a condition, because a pre-registered look has no selection among results to
    deflate.

    The recorded columns are the lab's history and are never written (design §1, plan invariant
    3). This is the *read*: what the same trial is judged as today, by the bars the lab holds
    today, rather than by the bars in force on its run date. That is R2's comparability, obtained
    without rewriting anything.

    **Nothing is ever returned verbatim.** Every one of the six conditions is decided here:

    - the four **threshold** owner conditions are re-derived from this row's recorded columns by
      ``owner_failures``, against the live ``tuning.MAX_DRAWDOWN`` / ``tuning.MIN_PROFIT_FACTOR``
      / ``dev._MIN_TRADES``, so a trial recorded under the old 15% drawdown bar reads against
      today's 20% one;
    - ``owner inputs`` is carried from the recorded string, because it is not a threshold and no
      constant re-decides it;
    - the **luck** test is decided on ``dsr_at(conn, trial, g.n)`` -- this trial's DSR **at the
      gate's current N**, never at the N it happened to be run under.

    **The luck test is always evaluated at the current N, and that is the whole of R2.** An
    earlier draft re-thresholded the recorded ``dsr`` only when ``n_trials_at_run`` already
    equalled the gate's N, and returned the row verbatim otherwise. That rule is **superseded and
    must not be implemented**: it would admit ``M0007-N20-RAW`` on a DSR of 0.9138 computed at
    N = 85 while judging ``M0022``'s variants on DSRs computed at N = 110 -- a candidate admitted
    for having been tried *earlier*, which is precisely the leaderboard-mixes-bars defect R2
    names and precisely the self-deception the lab exists to prevent. Re-evaluated at today's
    N = 110, ``M0007-N20-RAW`` is **0.8985** and does not clear 0.90.

    **A DSR that cannot be evaluated fails the luck test.** ``dsr_at`` returns None for the 54
    P7a seed rows, whose ``dsr`` is NULL by construction, and ``verdict`` then appends the luck
    label -- the same rule ``runner.trial_rows`` applies to a new trial
    (``if dsr is None or dsr < DSR_MIN``). Without it, ``F9-SPY200M70-MOM30`` (19.2% drawdown,
    MAR 0.64, no recorded DSR) would become eligible the moment the drawdown bar moved, on the
    strength of a luck test nobody ever ran. ``derived`` is False exactly in this case, and
    ``dsr`` is None with it.

    ``at`` resolves the gate once for a caller judging many trials. ``policy`` names a different
    policy for a read-only comparison (phase 5's ``lab luck``); the write paths never pass it.
    """
    g = gate(conn, policy) if at is None else at
    dsr = dsr_at(conn, trial, g.n)
    # Four re-derived from this row's columns against the live constants, one (`owner inputs`)
    # carried from the recorded string. Never parsed out of `failed` -- that string names the
    # bars in force on the run date, which are not today's.
    failed = owner_failures(trial)
    # The same rule runner.trial_rows applies to a new trial: a DSR that is None is a luck test
    # that was not passed. Nothing is admitted for being unmeasurable.
    if dsr is None or dsr < DSR_MIN:
        failed = failed + (DSR_LABEL,)
    return Verdict(
        dsr=dsr, failed=failed, eligible=not failed, n=g.n, policy=g.policy,
        derived=dsr is not None,
    )


def published_verdict(
    conn: sqlite3.Connection, trial: Mapping[str, Any] | sqlite3.Row, *, at: Gate | None = None
) -> Verdict:
    """``trial``'s verdict for a reader of the snapshot: ``verdict`` for a dev row, and the
    thresholds alone for a test row.

    ``verdict`` is defined on dev rows, because the luck test is a statement about **selection**
    among dev results. A test-window row has no selection to deflate -- that is the whole point
    of pre-registering one variant and looking once -- so its recorded ``dsr`` is a measurement,
    not a condition, and appending a luck label to it would invent a hurdle the lab never set.

    What a test row does still have is the four owner thresholds, and those are constants the
    owner moves: a test row recorded under the 15% drawdown bar must read against today's 20%
    one exactly as a dev row does. So ``owner_failures`` is applied to it (four re-derived from
    its own columns, ``owner inputs`` carried) and nothing else is.

    ``dsr`` carries the recorded number for a test row and the re-evaluated one for a dev row;
    ``derived`` says which, so a caller never has to guess. ``n`` and ``policy`` describe the
    gate either way, because they describe the lab, not the row.
    """
    g = gate(conn) if at is None else at
    if str(trial["window"]) == "dev":
        return verdict(conn, trial, at=g)
    failed = owner_failures(trial)
    dsr = None if trial["dsr"] is None else float(trial["dsr"])
    return Verdict(
        dsr=dsr, failed=failed, eligible=not failed, n=g.n, policy=g.policy, derived=False
    )


def luck_gated(trial: Mapping[str, Any] | sqlite3.Row) -> bool:
    """Does the luck gate apply to ``trial``? True for a dev row, False for a test look.

    The same split :func:`published_verdict` makes, named once so that no reader downstream has to
    re-derive it. A dev row is one result selected from many, so its Sharpe is deflated by the
    number of looks; a test-window row is a single pre-registered confirmatory look with no
    selection to deflate, so its recorded ``dsr`` is a measurement and not a condition.

    **This is not ``Verdict.derived``.** ``derived`` is False for a test row *and* for a dev row
    whose DSR could not be evaluated -- and those two are opposites. The first is a hurdle that
    does not apply; the second is a hurdle that was missed because nothing could be measured
    (the 54 P7a seed rows, whose ``dsr`` is NULL by construction). A reader that cannot tell them
    apart prints a pass where there is none: ``web/lib/sera/derive.conditionOk`` returned a green
    tick on ``M0021-B70-RAW``'s 0.513013 against a published bar of 0.90, which is R7.

    Published as ``trials[].luckGated`` from snapshot v4 so the web reads this answer rather than
    guessing at it. ``test_the_marker_names_the_same_split_published_verdict_makes`` holds this
    function and ``published_verdict`` to one rule, so the two cannot drift apart.
    """
    return str(trial["window"]) == "dev"


# The first words of the analysis section ``reevaluate_method`` appends. Not an idempotence key:
# the edge it guards can be taken at most once, because it leads out of the only status it may
# be taken from.
REEVALUATION_MARKER = "Re-evaluated under the "


@dataclass(frozen=True)
class Reevaluation:
    """What ``reevaluate_method`` found, and whether it moved the method."""

    method_id: str
    status_before: str
    status_after: str
    moved: bool
    gate: Gate
    unblocked: tuple[str, ...]  # candidate ids now eligible that the recorded column rejects
    derived: int                # dev trials whose DSR was evaluated at the gate's N
    unjudgeable: int            # dev trials with no evaluable DSR -- they fail the luck test
    dev_trials: int


def _blocking(trial: Mapping[str, Any] | sqlite3.Row) -> tuple[str, ...]:
    """The second, independent no: why this trial may **not** take the ``rejected`` edge.

    ``owner_failures`` derives the same four conditions and ``verdict`` reads it, so these
    comparisons are written here a second time **on purpose**. Two owner-set bars moved in this
    plan set -- the luck threshold (phase 4) and the drawdown threshold (phase 8) -- and a gate
    that loosens on two axes at once should not be able to promote a method through a single
    expression. A future change to ``owner_failures`` has to get past this too.

    Phrased as sentences with the numbers in them, because this text goes into the refusal a
    human reads.
    """
    from seer_engine.backtest import dev, tuning

    out: list[str] = []
    total, bench = trial["total_return"], trial["spy_tr_return"]
    if total is None or bench is None or float(total) <= float(bench):
        out.append(f"total return {total!r} does not beat SPY TR {bench!r}")
    dd = trial["max_drawdown"]
    if dd is None or float(dd) > tuning.MAX_DRAWDOWN:
        out.append(f"max drawdown {dd!r} is outside the {tuning.MAX_DRAWDOWN:.0%} bar")
    factor = trial["profit_factor"]
    if factor is None or float(factor) < tuning.MIN_PROFIT_FACTOR:
        out.append(f"profit factor {factor!r} is under {tuning.MIN_PROFIT_FACTOR}")
    count = trial["trades"]
    if count is None or int(count) < dev._MIN_TRADES:
        out.append(f"{count!r} closed trades is under {dev._MIN_TRADES}")
    if OWNER_INPUTS_LABEL in recorded_labels(trial["failed"]):
        out.append(
            f"it recorded {OWNER_INPUTS_LABEL!r}, which is a property of the candidate and which "
            f"no threshold re-decides"
        )
    return tuple(out)


def reevaluate_method(conn: sqlite3.Connection, method_id: str) -> Reevaluation:
    """Re-judge one method's dev trials under ``DSR_MIN`` and ``DSR_POLICY`` and, if that
    unblocks it, move it ``rejected -> dev-eligible``.

    **This is the only write path in the lab that reconsiders a verdict, and it is deliberately
    narrow.** It may take exactly one edge, from exactly one status, for exactly one reason:

    - the method must read ``rejected`` now. Any other status is left alone and ``moved`` is
      False; there is no other edge into ``dev-eligible`` from here, so a method already moved
      is not moved twice.
    - at least one of its dev trials must be eligible under the derived verdict and not already
      recorded eligible.
    - **every trial it relies on must clear all five conditions on a second, independently
      written check.** ``verdict`` already derives them; ``_blocking`` repeats the four
      comparisons against the same live constants in its own code, and refuses outright if the
      recorded ``failed`` carries ``OWNER_INPUTS_LABEL``. The repetition is the point: **two
      owner-set bars moved in this set** (the luck threshold here, the drawdown threshold in
      phase 8), and a gate that loosens on two axes at once is worth two noes written in two
      places. A trial that reads derived-eligible but does not clear this check is a
      contradiction inside this module, and it raises ``LabError`` rather than promoting quietly.

      Note what this check is **not**: it is not "the recorded failure was the luck label alone".
      That was the earlier draft's rule, and it is wrong now -- ``M0020-W-NOSTOP``'s recorded
      failure is ``"max DD <= 15%; DSR >= 0.95"`` and it *should* become eligible, because the
      owner moved both of those bars. The rule is about the numbers, not about the string.

    Nothing in ``trials`` is written: no row is inserted, updated or deleted, so the lab's N does
    not move and ``test_looks`` is untouched. What is written is one dated section appended to
    ``analysis`` (append-only by trigger, the shape ``record_promotion`` uses) naming the
    threshold, the policy, the N and the date that re-judged it, and the ``status`` column. Both
    happen in the caller's transaction, so a crash leaves neither.

    Neither the threshold nor the policy is a parameter. A write path that could unblock a method
    under any bar on request would make both constants decorative; the read-only comparison
    across policies is phase 5's ``lab luck``.

    The caller holds the transaction (``begin_immediate`` / ``with conn``), as every other
    writer in this module does.
    """
    from seer_engine.backtest import tuning  # for the drawdown bar named in the analysis text

    row = get_method(conn, method_id)
    if row is None:
        raise LabError(f"no method {method_id}")
    status = str(row["status"])
    trials = conn.execute(
        "SELECT * FROM trials WHERE method_id = ? AND window = 'dev' ORDER BY n", (method_id,)
    ).fetchall()
    g = gate(conn)
    unblocked: list[str] = []
    verdicts: dict[str, Verdict] = {}
    derived = 0
    for t in trials:
        v = verdict(conn, t, at=g)
        derived += 1 if v.derived else 0
        if not v.eligible:
            continue
        blocking = _blocking(t)
        if blocking:
            raise LabError(
                f"{t['candidate_id']} reads eligible at {DSR_LABEL} under the {g.policy} policy "
                f"(N={g.n}), but a second, independent check of its recorded numbers says "
                f"otherwise: {'; '.join(blocking)}. Refusing to move {method_id} -- a method may "
                f"only cross this edge when every one of the five conditions clears the bar in "
                f"force today."
            )
        if not bool(t["eligible"]):
            unblocked.append(str(t["candidate_id"]))
            verdicts[str(t["candidate_id"])] = v
    found = Reevaluation(
        method_id=method_id,
        status_before=status,
        status_after=status,
        moved=False,
        gate=g,
        unblocked=tuple(unblocked),
        derived=derived,
        unjudgeable=len(trials) - derived,
        dev_trials=len(trials),
    )
    if status != "rejected" or not unblocked:
        return found

    lines = [
        "# Re-evaluation",
        "",
        f"{REEVALUATION_MARKER}`{g.policy}` N policy at N = {g.n}, against `{DSR_LABEL}` and "
        f"a max drawdown bar of {tuning.MAX_DRAWDOWN:.0%} (both the owner's risk appetite, set "
        f"2026-10-07; LAB_LUCK_GATE_PLAN.md Decisions D1 and D6).",
        "",
        "Every condition was re-read against the bars in force today: the four threshold "
        "conditions from this method's own recorded columns, and the luck test from the same "
        "deflated Sharpe at the same N. No recorded column changed: `trials` is append-only, and "
        "every `dsr`, `eligible`, `failed` and `n_trials_at_run` this method recorded still reads "
        "exactly as it did -- including the `DSR >= 0.95` and `max DD <= 15%` labels that "
        "rejected it, which name the bars of their own day.",
        "",
        f"The {'variant' if len(unblocked) == 1 else 'variants'} this unblocks:",
        "",
    ]
    for cid in unblocked:
        v = verdicts[cid]
        shown = "None" if v.dsr is None else f"{v.dsr:.4f}"
        lines.append(f"- `{cid}`: DSR {shown} at N = {g.n}, against {DSR_LABEL}.")
    lines += [
        "",
        "The *set* of five P7a D8 conditions did not change, and `owner inputs` -- the one that "
        "is a property of the candidate rather than of a number -- was carried from the record "
        "untouched. Status moves `rejected` -> `dev-eligible`.",
    ]
    append_analysis(conn, method_id, "\n".join(lines))
    update_method(conn, method_id, status="dev-eligible")
    return Reevaluation(
        method_id=method_id,
        status_before=status,
        status_after="dev-eligible",
        moved=True,
        gate=g,
        unblocked=tuple(unblocked),
        derived=derived,
        unjudgeable=len(trials) - derived,
        dev_trials=len(trials),
    )


def reevaluate(
    conn: sqlite3.Connection, method_ids: Sequence[str] | None = None
) -> list[Reevaluation]:
    """``reevaluate_method`` over ``method_ids``, or over every ``rejected`` method when None.

    The caller holds the transaction, so the whole sweep is one atomic unit: either every method
    it unblocks moves, or none does.
    """
    if method_ids is None:
        method_ids = [
            str(r[0])
            for r in conn.execute("SELECT id FROM methods WHERE status = 'rejected' ORDER BY id")
        ]
    return [reevaluate_method(conn, m) for m in method_ids]


# --------------------------------------------------------------------------- trial moments


@dataclass(frozen=True)
class MomentsRow:
    """The deflated Sharpe's inputs for one dev trial, as they were at the moment it was judged.

    ``dev.deflated_sharpe(sr_daily, n_at_run, var_trials, t, skew, kurt)`` reproduces that
    trial's recorded ``dsr`` exactly; substituting another N re-judges it under another policy
    without re-running the backtest. That is the whole point of the table: ``trials`` is
    append-only and its ``dsr`` is frozen at the N of its run date, so a verdict can only be
    recomputed from inputs that were kept.

    - ``trial_n``   the ``trials.n`` this describes. ``insert_trials`` assigns it, so a row built
                    before the insert carries ``0`` and is stamped with
                    ``dataclasses.replace(row, trial_n=n)`` afterwards; ``insert_moments``
                    refuses ``0``.
    - ``sr_daily``  the daily Sharpe (``daily_moments(...)[0]``), not annualized.
    - ``t``         the number of daily returns the trial produced.
    - ``skew``      skewness of those returns (``m3 / m2**1.5``).
    - ``kurt``      kurtosis of those returns, **not** excess (3.0 for a normal).
    - ``var_trials``the variance of the daily Sharpe across the trials this one was deflated
                    against, or None when the lab had fewer than two -- the same condition under
                    which ``trials.dsr`` is NULL.
    - ``n_at_run``  the N used, which equals the trial's ``n_trials_at_run``.
    - ``measured``  an ISO timestamp: for a trial recorded by ``lab run`` this is the trial's own
                    ``run_at``, so the pair is one measurement with one stamp.
    """

    trial_n: int
    sr_daily: float
    t: int
    skew: float
    kurt: float
    var_trials: float | None
    n_at_run: int
    measured: str


MOMENTS_COLUMNS: tuple[str, ...] = tuple(f.name for f in fields(MomentsRow))


def insert_moments(conn: sqlite3.Connection, rows: Sequence[MomentsRow]) -> None:
    """Record the DSR's inputs for trials that already exist (the caller holds the transaction).

    Append-only and one row per trial: a trial that already has moments is refused rather than
    overwritten, so a backfill (``lab remeasure``) can be made idempotent by skipping what
    ``moments_of`` already returns instead of racing the triggers.

    ``trial_n`` must be a real trial number -- the foreign key enforces that the trial exists,
    and the explicit check below turns the pre-insert sentinel ``0`` into a readable refusal
    rather than a foreign-key error from inside a long run.
    """
    cols = ", ".join(f'"{c}"' for c in MOMENTS_COLUMNS)
    marks = ", ".join("?" for _ in MOMENTS_COLUMNS)
    for r in rows:
        if r.trial_n <= 0:
            raise LabError(
                f"moments need the trial number insert_trials assigned, got {r.trial_n!r}: "
                "insert the trial first, then stamp its moments with dataclasses.replace"
            )
        if moments_of(conn, r.trial_n) is not None:
            raise LabError(
                f"trial {r.trial_n} already has recorded moments: trial_moments is append-only, "
                "and a second measurement of the same trial would be a second verdict"
            )
        conn.execute(
            f"INSERT INTO trial_moments ({cols}) VALUES ({marks})",
            [getattr(r, c) for c in MOMENTS_COLUMNS],
        )


def moments_of(conn: sqlite3.Connection, trial_n: int) -> sqlite3.Row | None:
    """The recorded DSR inputs for one trial, or None when it has none.

    None is the normal answer for every trial recorded before this table existed, and for any
    trial whose daily returns had no computable moments (fewer than two returns, or zero
    variance) -- which is exactly the set of trials whose ``dsr`` is NULL. A caller that
    re-evaluates must therefore have a fallback for None; it is never an error.

    A database on schema v1 or v2 has no ``trial_moments`` table at all, which is the limiting
    case of "recorded before this table existed" and is answered the same way. ``connect``
    migrates, so this can only be a ``connect_readonly`` caller -- ``snapshot`` is one, and its
    promise to work on an unmigrated read-only connection is what this branch keeps.
    """
    if conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'trial_moments'"
    ).fetchone() is None:
        return None
    return conn.execute(
        "SELECT * FROM trial_moments WHERE trial_n = ?", (int(trial_n),)
    ).fetchone()


# --------------------------------------------------------------------------- ideas_seen


def mark_seen(conn: sqlite3.Connection, key: str, method_id: str | None = None, note: str = "") -> bool:
    """Record a dedupe key; False when it was already there."""
    key = key.strip()
    if ":" not in key:
        raise LabError(f"a seen key looks like kind:value (url:…, concept:…), got {key!r}")
    cur = conn.execute(
        "INSERT OR IGNORE INTO ideas_seen (key, method_id, note, added) VALUES (?, ?, ?, ?)",
        (key, method_id, note, now_iso()),
    )
    return cur.rowcount == 1


def seen(conn: sqlite3.Connection, pattern: str) -> list[sqlite3.Row]:
    """Seen keys containing ``pattern`` (case-insensitive)."""
    return conn.execute(
        "SELECT * FROM ideas_seen WHERE lower(key) LIKE ? OR lower(note) LIKE ? ORDER BY key",
        (f"%{pattern.lower()}%", f"%{pattern.lower()}%"),
    ).fetchall()


# --------------------------------------------------------------------------- insights


def add_insight(conn: sqlite3.Connection, *, kind: str, title: str, body: str,
                method_id: str | None = None) -> int:
    """Append one entry to the lab journal: an observation across methods, a hypothesis worth
    testing, data or a feature the lab lacks, a risk, or a batch synthesis (the state of the
    search after a batch of methods). The food-for-thought record, shown at seertrade.site/sera."""
    if kind not in INSIGHT_KINDS:
        raise LabError(f"insight kind must be one of {INSIGHT_KINDS}, got {kind!r}")
    if method_id is not None and get_method(conn, method_id) is None:
        raise LabError(f"no method {method_id}")
    cur = conn.execute(
        "INSERT INTO insights (kind, title, body, method_id, added) VALUES (?, ?, ?, ?, ?)",
        (kind, title.strip(), body.strip(), method_id, now_iso()),
    )
    return int(cur.lastrowid)


# --------------------------------------------------------------------------- export


def _sheet_rows(conn: sqlite3.Connection, sql: str) -> tuple[list[str], list[list[Any]]]:
    cur = conn.execute(sql)
    header = [d[0] for d in cur.description]
    return header, [list(r) for r in cur.fetchall()]


LEADERBOARD_SQL = """
SELECT t.n, t.candidate_id, t.method_id, m.name AS method, m.family, t.window, t.start, t."end",
       t.total_return, t.spy_tr_return, t.cagr, t.max_drawdown, t.profit_factor, t.trades,
       t.sharpe, t.mar, t.dsr, t.n_trials_at_run, t.eligible, t.failed
FROM trials t JOIN methods m ON m.id = t.method_id
ORDER BY t.mar IS NULL, t.mar DESC, t.n
"""


def export_xlsx(conn: sqlite3.Connection, path: Path = XLSX_PATH) -> Path:
    """Write the lab as a workbook: summary, leaderboard, methods, trials (without curves), ideas_seen."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    for row in summary_rows(conn):
        summary.append(list(row))
    trial_cols = ", ".join(f't."{c}"' for c in ("n",) + TRIAL_COLUMNS if c != "curve_json")
    sheets = (
        ("leaderboard", LEADERBOARD_SQL),
        ("methods", "SELECT * FROM methods ORDER BY id"),
        ("trials", f"SELECT {trial_cols} FROM trials t ORDER BY t.n"),
        ("ideas_seen", "SELECT * FROM ideas_seen ORDER BY key"),
        ("insights", "SELECT * FROM insights ORDER BY id"),
    )
    for title, sql in sheets:
        ws = wb.create_sheet(title)
        header, rows = _sheet_rows(conn, sql)
        ws.append(header)
        for r in rows:
            ws.append([None if isinstance(v, float) and not math.isfinite(v) else v for v in r])
        ws.freeze_panes = "A2"
        for i, name in enumerate(header, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, len(name) + 2))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def summary_rows(conn: sqlite3.Connection) -> list[tuple[str, Any]]:
    by_status = conn.execute("SELECT status, count(*) FROM methods GROUP BY status ORDER BY status").fetchall()
    out: list[tuple[str, Any]] = [
        ("dev trials (N for the DSR)", dev_trial_count(conn)),
        ("test-window looks used", test_looks(conn)),
        ("methods", int(conn.execute("SELECT count(*) FROM methods").fetchone()[0])),
        ("insights", int(conn.execute("SELECT count(*) FROM insights").fetchone()[0])),
    ]
    out += [(f"methods {r[0]}", r[1]) for r in by_status]
    return out


def curve_json(curve: Iterable[tuple[date, float]]) -> str:
    return json.dumps([[d.isoformat(), round(v, 6)] for d, v in curve], separators=(",", ":"))


def as_mapping(row: sqlite3.Row) -> Mapping[str, Any]:
    return {k: row[k] for k in row.keys()}


# --------------------------------------------------------------------------- web snapshot

# 2 adds the derived verdict to every trial (``failedNow`` / ``eligibleNow`` / ``dsrNow``). The
# bump is not cosmetic: before it, the snapshot published the gate's bars live and every trial's
# verdict as recorded history, so a page drawing both read `0.912 < 0.90` off one trial -- the
# tick from the row's recorded `DSR >= 0.95`, the number from the live bar beside it.
#
# 3 adds ``paper``: which roster entry is which lab method (``paper.roster.LAB_PROVENANCE``). The
# web had no way to answer that -- ``rules_id`` is shared by a dozen methods, so it is not a key --
# and a map hand-written in the site would drift the next promotion night in silence.
#
# 4 adds ``luckGated``: does the luck gate apply to this row at all. The lab produced its first two
# test-window looks on 2026-10-07, and ``published_verdict`` rightly does not luck-gate them -- a
# pre-registered confirmatory look has no selection to deflate. Nothing said so, so every consumer
# re-derived the rule and the web derived it wrongly: a score of 0.513013 rendered as a cleared
# hurdle against a published bar of 0.90. The marker is a third state, not a second one: it
# separates "this hurdle does not apply" from "this hurdle could not be measured, so it was
# missed", which ``Verdict.derived`` conflates (R7).
SNAPSHOT_VERSION = 4
EXPECTED_FAILURE_SEP = "\n\nExpected failure: "  # how lab ideas write their hypothesis


def snapshot_path(db: Path) -> Path:
    """The web snapshot of a database: ``<repo>/web/data/lab.json`` for ``<repo>/lab/lab.sqlite``.

    Derived from the database's own checkout, not from this code's: a worktree session with
    ``SEER_LAB_DB`` pointing at the main checkout writes the snapshot into that checkout, next to
    the database it exports, and ``lab stage`` stages both there."""
    return Path(db).resolve().parent.parent / "web" / "data" / "lab.json"


def _num(x: Any) -> float | None:
    """A JSON-safe number: None and non-finite values become None, the rest rounds to 6 dp."""
    if x is None:
        return None
    x = float(x)
    return round(x, 6) if math.isfinite(x) else None


def _dicts(conn: sqlite3.Connection, sql: str) -> list[dict[str, Any]]:
    """Rows as dicts, whatever the connection's row_factory."""
    cur = conn.execute(sql)
    names = [d[0] for d in cur.description]
    return [dict(zip(names, row)) for row in cur.fetchall()]


def _snapshot_method(m: Mapping[str, Any]) -> dict[str, Any]:
    hypothesis, sep, expected = m["hypothesis"].partition(EXPECTED_FAILURE_SEP)
    return {
        "id": m["id"],
        "name": m["name"],
        "family": m["family"],
        "parentId": m["parent_id"],
        "sourceKind": m["source_kind"],
        "sourceRef": m["source_ref"],
        "hypothesis": hypothesis,
        "expectedFailure": expected if sep else None,
        "status": m["status"],
        "analysis": m["analysis"],
        "verdict": m["verdict"],
        "blockedOn": m["blocked_on"],
        "created": m["created"],
        "updated": m["updated"],
        "historical": m["id"].startswith("H-"),
    }


def _snapshot_curve(text: str) -> list[list[Any]]:
    out: list[list[Any]] = []
    for day, value in json.loads(text):
        v = _num(value)
        if v is not None:
            out.append([str(day), v])
    return out


def _snapshot_trial(t: Mapping[str, Any], v: Verdict) -> dict[str, Any]:
    """One ``trials`` row as the web reads it: the record, **and** the verdict it reads as now.

    Both, and labelled as such, because they answer different questions and the site asks both.
    ``failed`` / ``eligible`` / ``dsr`` / ``nTrialsAtRun`` are the row as the lab wrote it on its
    run date -- append-only history, and the technical record a reader reruns from. ``failedNow``
    / ``eligibleNow`` / ``dsrNow`` are ``published_verdict``: the same row judged by the bars in
    force at export time, at the gate's current N.

    Publishing only the first is what made ``/sera/methods/M0022`` print ``0.912 < 0.90``: the
    gate block beside it is resolved live (``_gate_n``, ``DSR_MIN``), so a page that took its
    ticks from ``failed`` and its numbers from ``gate`` was reading two different days at once.
    A page must take **both** from the ``Now`` fields; ``failed`` is for the record panel only.

    ``luckGated`` says whether the luck hurdle applies to this row at all (:func:`luck_gated`).
    Without it the verdict is ambiguous in exactly one place: a test look's ``failedNow`` omits
    the luck label because the gate does not apply, and a reader with only ``failedNow`` to go on
    cannot tell that from a hurdle that was cleared. It published a tick on 0.513013 against a
    bar of 0.90. **Three states, not two**: gated and cleared, gated and missed (``dsrNow`` may be
    null -- nothing is admitted for being unmeasurable), and not gated at all.
    """
    pf = t["profit_factor"]
    return {
        "n": int(t["n"]),
        "methodId": t["method_id"],
        "candidateId": t["candidate_id"],
        "rulesId": t["rules_id"],
        "allocatorId": t["allocator_id"],
        "configText": t["config_text"],
        "window": t["window"],
        "start": t["start"],
        "end": t["end"],
        "gitSha": t["git_sha"],
        "runAt": t["run_at"],
        "totalReturn": _num(t["total_return"]),
        "cagr": _num(t["cagr"]),
        "maxDrawdown": _num(t["max_drawdown"]),
        "profitFactor": _num(pf),
        "pfInfinite": pf is not None and float(pf) == math.inf,
        "trades": int(t["trades"]),
        "sharpe": _num(t["sharpe"]),
        "exposure": _num(t["exposure"]),
        "turnover": _num(t["turnover"]),
        "worstYear": None if t["worst_year"] is None else int(t["worst_year"]),
        "worstYearReturn": _num(t["worst_year_return"]),
        "spyTrReturn": _num(t["spy_tr_return"]),
        "spyTrCagr": _num(t["spy_tr_cagr"]),
        "mar": _num(t["mar"]),
        # The record: what the lab said on the run date, by the bars of that day.
        "failed": [f for f in t["failed"].split("; ") if f],
        "eligible": bool(t["eligible"]),
        "dsr": _num(t["dsr"]),
        "nTrialsAtRun": int(t["n_trials_at_run"]),
        # Whether the luck hurdle applies to this row at all. Published so the web never has to
        # re-derive the dev/test rule -- which it did, and got wrong.
        "luckGated": luck_gated(t),
        # The verdict: the same row by the bars in force now, at the gate's N (gate.dsrN).
        "failedNow": list(v.failed),
        "eligibleNow": v.eligible,
        "dsrNow": _num(v.dsr),
        "curve": _snapshot_curve(t["curve_json"]),
    }


def _gate_n(conn: sqlite3.Connection) -> tuple[dict[str, Any], Gate]:
    """The multiple-testing N the luck gate deflates by, and the evidence behind it.

    Returns the published block **and the same N as a ``Gate``**, from one resolution, because
    ``snapshot`` needs both: the block it publishes and the gate it judges every trial's
    ``failedNow`` against. Two calls would be two eigendecompositions of 110 curves, and -- worse
    than slow -- two places the published N could come from.

    Three keys rather than one number, because the number alone is not reviewable. A reader of
    ``web/data/lab.json`` -- or of seertrade.site/sera, which draws this beside the luck bar --
    has to be able to see *which* policy produced the N and *what measurement* that policy rests
    on. That matters more here than it would if the policy had changed: it did not. ``DSR_POLICY``
    ships as ``all-trials`` on purpose (design §7.2), and publishing the name, the count and the
    evidence is what keeps the lever that was deliberately not pulled visible on the site rather
    than buried in a plan file.

    Resolved at snapshot time from ``DSR_POLICY`` and never stored, so flipping that one constant
    moves the published gate on the next export with no data change -- the same property the
    derived verdict has. Reads only ``trials``, which schema v1 and v2 share, so ``snapshot``'s
    promise to work on a read-only, unmigrated connection still holds.
    """
    from seer_engine.lab import npolicy

    n = npolicy.effective_n(conn, DSR_POLICY)
    return {"dsrPolicy": n.policy, "dsrN": n.n, "dsrNBasis": n.basis}, Gate(
        n=int(n.n), policy=str(n.policy)
    )


def snapshot(conn: sqlite3.Connection) -> dict[str, Any]:
    """The whole lab as the web's ``LabSnapshot`` (SERA_LAB_SITE_PLAN.md Interface Contract).

    Deterministic: the same database gives the same value; there is no wall-clock time in it
    (``asOf`` is the latest timestamp found in the data, "" for an empty lab). Reads only what
    schema v1 and v2 share and never touches ``meta``, so it works on a read-only connection to
    a database that has not been migrated. Gate and data facts come from the engine's constants;
    the gate's ``dsrPolicy`` / ``dsrN`` / ``dsrNBasis`` are resolved from ``DSR_POLICY`` against
    ``trials`` at export time rather than stored, so the published gate follows the constant
    (design §7).

    **Every trial is published twice over: as recorded, and as judged now.** The gate block has
    always been resolved at export time; from v2 each trial's ``failedNow`` / ``eligibleNow`` /
    ``dsrNow`` is resolved against that same gate, in the same pass, by ``published_verdict``.
    The two have to move together or not at all -- a snapshot that published live bars beside
    recorded verdicts is what printed ``0.912 < 0.90`` on the site.

    ``paper`` is the one block that is not a read of this database: it is
    ``paper.roster.LAB_PROVENANCE``, which names the lab method and variant behind every roster
    entry. It travels here because the web has no other way to get it -- ``rules_id`` is shared by
    a dozen methods, so it cannot be the key -- and because the alternative, a map written by hand
    in the site, goes stale on the next promotion night without anything failing. ``reason`` is
    deliberately left out: the roster's prose belongs on the roster, not in the lab's snapshot.
    """
    from seer_engine import dates, research
    from seer_engine.backtest import dev, tuning
    from seer_engine.lab import seed
    from seer_engine.paper import roster

    def count(sql: str) -> int:
        return int(conn.execute(sql).fetchone()[0])

    by_status: dict[str, int] = dict.fromkeys(STATUSES, 0)
    for status, n in conn.execute("SELECT status, count(*) FROM methods GROUP BY status ORDER BY status"):
        by_status[str(status)] = int(n)
    as_of = conn.execute(
        "SELECT max(ts) FROM (SELECT max(updated) AS ts FROM methods UNION ALL "
        "SELECT max(run_at) FROM trials UNION ALL SELECT max(added) FROM insights UNION ALL "
        "SELECT max(added) FROM ideas_seen)"
    ).fetchone()[0]
    bench = seed.benchmark_curves()
    gate_block, g = _gate_n(conn)
    return {
        "version": SNAPSHOT_VERSION,
        "asOf": as_of or "",
        "gate": {
            "maxDrawdown": tuning.MAX_DRAWDOWN,
            "minProfitFactor": tuning.MIN_PROFIT_FACTOR,
            "minTrades": dev._MIN_TRADES,
            "dsrMin": DSR_MIN,
            **gate_block,
            "devStart": research.STORE_START.isoformat(),
            "devEnd": dev.DEV_END.isoformat(),
            "testStart": dates.next_session(dev.DEV_END).isoformat(),
        },
        "data": {
            "storeStart": research.STORE_START.isoformat(),
            "membershipStart": dev.MEMBERSHIP_START.isoformat(),
            "fxStart": dev.FX_START.isoformat(),
            "fingerprints": [
                str(r[0]) for r in conn.execute(
                    "SELECT DISTINCT store_fingerprint FROM trials ORDER BY store_fingerprint"
                )
            ],
            "barRows": seed.P7A_BAR_ROWS,
            "symbolsRequested": seed.P7A_SYMBOLS_REQUESTED,
            "symbolsServed": seed.P7A_SYMBOLS_SERVED,
            "dividendRows": seed.P7A_DIVIDEND_ROWS,
        },
        "summary": {
            "devTrials": count("SELECT count(*) FROM trials WHERE window = 'dev'"),
            "testLooks": count("SELECT count(*) FROM trials WHERE window = 'test'"),
            "methods": count("SELECT count(*) FROM methods"),
            "labMethods": count("SELECT count(*) FROM methods WHERE id GLOB 'M*'"),
            "historicalMethods": count("SELECT count(*) FROM methods WHERE id GLOB 'H-*'"),
            "insights": count("SELECT count(*) FROM insights"),
            "byStatus": by_status,
        },
        "benchmark": {
            "spyTr": [[d, v] for d, v in bench["spy_tr"]],
            "spyPrice": [[d, v] for d, v in bench["spy_price"]],
        },
        "methods": [_snapshot_method(m) for m in _dicts(conn, "SELECT * FROM methods ORDER BY id")],
        # One gate, resolved once and passed down: `published_verdict` judges every row against
        # the same N, which is also the `gate.dsrN` published above.
        "trials": [
            _snapshot_trial(t, published_verdict(conn, t, at=g))
            for t in _dicts(conn, "SELECT * FROM trials ORDER BY n")
        ],
        "insights": [
            {"id": int(i["id"]), "kind": i["kind"], "title": i["title"], "body": i["body"],
             "methodId": i["method_id"], "added": i["added"]}
            for i in _dicts(conn, "SELECT * FROM insights ORDER BY id")
        ],
        "ideasSeen": [
            {"key": s["key"], "methodId": s["method_id"], "note": s["note"], "added": s["added"]}
            for s in _dicts(conn, "SELECT * FROM ideas_seen ORDER BY key")
        ],
        # Not from the database: the roster's own record of where each paper entry came from.
        "paper": [
            {"strategyId": sid, "methodId": p.method_id, "candidateId": p.candidate_id,
             "labStatus": p.lab_status, "basis": p.basis}
            for sid, p in sorted(roster.LAB_PROVENANCE.items())
        ],
    }


def snapshot_json(conn: sqlite3.Connection) -> str:
    """``snapshot`` as the exact bytes of ``web/data/lab.json``: compact, key order as built,
    UTF-8 kept, one trailing newline. ``allow_nan=False``: a stray NaN/inf fails loudly."""
    return json.dumps(snapshot(conn), ensure_ascii=False, separators=(",", ":"), allow_nan=False) + "\n"


def write_snapshot(conn: sqlite3.Connection, path: Path) -> bool:
    """Write ``snapshot_json(conn)`` to ``path`` (temp file + atomic replace). False when the
    file already held exactly that, in which case it is left untouched."""
    data = snapshot_json(conn).encode("utf-8")
    path = Path(path)
    if path.is_file() and path.read_bytes() == data:
        return False
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_bytes(data)
    os.replace(tmp, path)
    return True
