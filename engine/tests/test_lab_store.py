"""The lab database's rules (method lab design §1): append-only trials, forward-only status,
growing analysis that refuses an identical re-append, one test look per configuration, the seed
import and the export.

The duplicate-note refusal is tested from both ends here: at ``store.append_analysis``, and
through ``lab note`` itself, because the bug it exists to stop was half a store bug and half a
silent command."""

from __future__ import annotations

import argparse
import logging
import sqlite3
from dataclasses import replace
from datetime import date
from decimal import Decimal

import pytest

from seer_engine.backtest.registry import REGISTRY
from seer_engine.commands import lab as lab_cmd
from seer_engine.lab import store
from seer_engine.lab import seed as seed_mod
from seer_engine.lab.method import config_digest
from seer_engine.lab.seed import seed


@pytest.fixture()
def conn(tmp_path):
    c = store.connect(tmp_path / "lab.sqlite")
    yield c
    c.close()


def _method(conn, mid="M0001", status="idea"):
    with conn:
        store.add_method(conn, id=mid, name="n", family="f", source_kind="knowledge", hypothesis="h", status=status)


def _trial(**kw) -> store.TrialRow:
    base = dict(
        method_id="M0001", candidate_id="M0001-A", config_digest="d1", config_text="t", rules_id="r",
        allocator_id="a", window="dev", start="2000-01-03", end="2015-10-16", store_fingerprint="fp",
        git_sha="abc", run_at="2026-10-04T00:00:00+00:00", total_return=1.0, cagr=0.1, max_drawdown=0.2,
        profit_factor=1.5, trades=200, sharpe=0.8, exposure=0.9, turnover=1.0, worst_year=2008,
        worst_year_return=-0.2, spy_tr_return=0.5, spy_tr_cagr=0.07, mar=0.5, failed="max DD <= 15%",
        eligible=False, dsr=0.5, n_trials_at_run=55, curve_json="[]",
    )
    base.update(kw)
    return store.TrialRow(**base)


def test_trials_are_append_only(conn):
    _method(conn)
    with conn:
        store.insert_trials(conn, [_trial()])
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        conn.execute("UPDATE trials SET mar = 9")
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        conn.execute("DELETE FROM trials")
    with pytest.raises(sqlite3.IntegrityError, match="never deleted"):
        conn.execute("DELETE FROM methods")


def test_a_configuration_runs_once_per_window(conn):
    _method(conn)
    with conn:
        store.insert_trials(conn, [_trial()])
    with pytest.raises(store.LabError, match="already has a dev trial"):
        store.insert_trials(conn, [_trial(candidate_id="M0001-B")])
    with conn:
        store.insert_trials(conn, [_trial(window="test")])
    with pytest.raises(store.LabError, match="already has a test trial"):
        store.insert_trials(conn, [_trial(window="test", candidate_id="M0001-C")])
    assert store.test_looks(conn) == 1
    assert store.dev_trial_count(conn) == 1


def test_best_dev_eligible_is_the_highest_mar_and_breaks_ties_on_the_trial_number(conn):
    # Since LAB_LUCK_GATE_PLAN.md phase 4, "eligible" here is `store.verdict(...).eligible`, not
    # the recorded `eligible` column -- so the rows have to be eligible in their NUMBERS, under
    # the bars in force now. Two consequences for this fixture, neither of which touches what the
    # test is about (MAR order, the tie on `n`, the NULL MAR and the test trial):
    #   * the Sharpes differ, so the lab has a trial-Sharpe variance for `dsr_at` to deflate by
    #     -- with one repeated Sharpe there is no dispersion and no DSR is evaluable at all;
    #   * `n_trials_at_run` is the 4 dev rows below, so each DSR is read back at its own N, and
    #     C is kept out by a 55% drawdown rather than by its recorded `failed` string.
    _method(conn)
    assert store.best_dev_eligible(conn, "M0001") is None
    with conn:
        store.insert_trials(conn, [
            _trial(candidate_id="M0001-A", config_digest="da", mar=0.9, eligible=True, failed="",
                   dsr=0.95, sharpe=0.80, n_trials_at_run=4),
            _trial(candidate_id="M0001-B", config_digest="db", mar=0.9, eligible=True, failed="",
                   dsr=0.95, sharpe=0.94, n_trials_at_run=4),
            _trial(candidate_id="M0001-C", config_digest="dc", mar=1.4, max_drawdown=0.55,
                   dsr=0.95, sharpe=1.10, n_trials_at_run=4),  # not eligible: drawdown
            _trial(candidate_id="M0001-D", config_digest="dd", mar=None, eligible=True, failed="",
                   dsr=0.95, sharpe=1.46, n_trials_at_run=4),
            _trial(candidate_id="M0001-E", config_digest="de", window="test", mar=2.0,
                   eligible=True, failed=""),
        ])
    # Judge at the N these rows were RECORDED at -- `n_trials_at_run=4`, which is the dev row
    # count -- rather than at whatever `store.DSR_POLICY` resolves to. That is what makes
    # `store.verdict`'s re-evaluation the identity here, which is the condition the comment above
    # describes and which every assertion below rests on. Derived from the data, not typed, and
    # it names no shipped constant: the subject is `best_dev_eligible`'s ordering, not the gate.
    at = store.Gate(n=store.dev_trial_count(conn), policy="all-trials")
    best = store.best_dev_eligible(conn, "M0001", at=at)
    assert best["candidate_id"] == "M0001-A"  # the tie breaks on n, and the test trial is not it
    assert store.best_dev_eligible(conn, "M0002", at=at) is None


def _moments(**kw) -> store.MomentsRow:
    base = dict(trial_n=1, sr_daily=0.0501, t=3959, skew=-0.31, kurt=7.2, var_trials=2.395e-04,
                n_at_run=110, measured="2026-10-07T00:00:00+00:00")
    base.update(kw)
    return store.MomentsRow(**base)


def test_trial_moments_are_append_only_and_at_most_one_per_trial(conn):
    _method(conn)
    with conn:
        ns = store.insert_trials(conn, [_trial()])
    assert store.moments_of(conn, ns[0]) is None  # a trial recorded before this table has none
    with conn:
        store.insert_moments(conn, [_moments(trial_n=ns[0])])
    got = store.moments_of(conn, ns[0])
    assert list(got.keys()) == list(store.MOMENTS_COLUMNS)
    assert (got["trial_n"], got["t"], got["n_at_run"]) == (ns[0], 3959, 110)
    assert got["sr_daily"] == 0.0501 and got["skew"] == -0.31 and got["kurt"] == 7.2
    assert got["var_trials"] == 2.395e-04 and got["measured"] == "2026-10-07T00:00:00+00:00"

    with pytest.raises(store.LabError, match="already has recorded moments"):
        store.insert_moments(conn, [_moments(trial_n=ns[0], t=10)])
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        conn.execute("UPDATE trial_moments SET t = 9")
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        conn.execute("DELETE FROM trial_moments")
    assert store.moments_of(conn, ns[0])["t"] == 3959


def test_moments_need_a_real_trial_number(conn):
    _method(conn)
    with conn:
        store.insert_trials(conn, [_trial()])
    with pytest.raises(store.LabError, match="insert_trials assigned"):
        store.insert_moments(conn, [_moments(trial_n=0)])
    with pytest.raises(sqlite3.IntegrityError, match="FOREIGN KEY"):
        with conn:
            store.insert_moments(conn, [_moments(trial_n=999)])
    assert store.moments_of(conn, 999) is None
    with conn:  # var_trials is nullable: a lab with fewer than two Sharpes has no variance
        store.insert_moments(conn, [_moments(trial_n=1, var_trials=None, n_at_run=1)])
    assert store.moments_of(conn, 1)["var_trials"] is None


def test_status_only_moves_forward(conn):
    # `rejected -> dev-eligible` is the one edge out of `rejected` (LAB_LUCK_GATE_PLAN.md
    # Decision D2) and is exercised in test_lab_gate_policy.py, where the Python guard that
    # admits it lives. Every other move out of `rejected` is still refused, which is what this
    # test is about.
    _method(conn)
    with conn:
        store.update_method(conn, "M0001", status="registered")
        store.update_method(conn, "M0001", status="rejected")
    with pytest.raises(store.LabError, match="forward"):
        store.update_method(conn, "M0001", status="promoted")
    with pytest.raises(store.LabError, match="forward"):
        store.update_method(conn, "M0001", status="paper")
    with pytest.raises(store.LabError, match="forward"):
        store.update_method(conn, "M0001", status="idea")


def test_a_new_method_starts_as_an_idea_or_registered(conn):
    with pytest.raises(store.LabError):
        store.add_method(conn, id="M0001", name="n", family="f", source_kind="knowledge",
                         hypothesis="h", status="dev-eligible")
    with pytest.raises(sqlite3.IntegrityError):
        store.add_method(conn, id="X1", name="n", family="f", source_kind="knowledge", hypothesis="h")


def test_append_analysis_does_not_double_the_authors_own_date_heading(conn):
    today = date.today().isoformat()
    _method(conn)
    with conn:
        store.append_analysis(conn, "M0001", f"### {today}\n\nthe note")
        store.append_analysis(conn, "M0001", "a note with no heading of its own")
        store.append_analysis(conn, "M0001", "### 2020-01-01\n\nsomeone else's date")
    text = store.get_method(conn, "M0001")["analysis"]
    assert f"### {today}\n\n### {today}" not in text
    assert text.count(f"### {today}") == 3  # one head per append, never two
    assert text.startswith(f"### {today}\n\nthe note")
    assert f"### {today}\n\n### 2020-01-01" in text  # a stale date is headed, not trusted


def test_analysis_grows_and_hypothesis_freezes(conn):
    _method(conn)
    with conn:
        store.append_analysis(conn, "M0001", "first")
        store.append_analysis(conn, "M0001", "second")
    text = store.get_method(conn, "M0001")["analysis"]
    assert text.index("first") < text.index("second")
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        conn.execute("UPDATE methods SET analysis = 'rewritten' WHERE id = 'M0001'")
    with conn:
        store.update_method(conn, "M0001", hypothesis="h2")  # still an idea
        store.update_method(conn, "M0001", status="registered")
    with pytest.raises(store.LabError, match="frozen"):
        store.update_method(conn, "M0001", hypothesis="h3")


def test_append_analysis_refuses_an_identical_re_append(conn):
    """The M0245 bug: the same body twice, minutes apart, because the terminal said it failed."""
    _method(conn)
    body = "The six arms all beat the placebo.\n\n- FRESH30 +1.5\n- STALE30 +3.9\n"
    with conn:
        store.append_analysis(conn, "M0001", body)
    before = store.get_method(conn, "M0001")["analysis"]
    with pytest.raises(store.LabError, match="already the newest section"):
        with conn:
            store.append_analysis(conn, "M0001", body)
    assert store.get_method(conn, "M0001")["analysis"] == before


def test_append_analysis_catches_a_retry_that_only_differs_in_whitespace(conn):
    """A re-save between two tries is not a new note; the stored bytes stay the author's own."""
    _method(conn)
    body = "A note with trailing spaces   \nand a second line\n"
    with conn:
        store.append_analysis(conn, "M0001", body)
    before = store.get_method(conn, "M0001")["analysis"]
    for retry in (
        "\n\n  " + body + "   \n\n",                          # surrounding whitespace
        body.replace("\n", "\r\n"),                           # CRLF endings
        "A note with trailing spaces\nand a second line",     # per-line trailing spaces gone
    ):
        with pytest.raises(store.LabError, match="nothing was appended"):
            with conn:
                store.append_analysis(conn, "M0001", retry)
    assert store.get_method(conn, "M0001")["analysis"] == before
    assert "trailing spaces   \n" in before  # normalised for the compare, never for the record


def test_append_analysis_still_takes_a_different_body_and_still_heads_it(conn):
    """Newest section only: a body that is no longer the newest one is a new act and lands."""
    today = date.today().isoformat()
    _method(conn)
    with conn:
        store.append_analysis(conn, "M0001", "first")
        store.append_analysis(conn, "M0001", "second")
        store.append_analysis(conn, "M0001", "first")
    text = store.get_method(conn, "M0001")["analysis"]
    assert text.count(f"### {today}") == 3
    assert f"### {today}\n\n### {today}" not in text
    assert text.index("second") < text.rindex("first")


def test_source_sha_is_set_once_and_later_shas_go_to_method_amendments(conn):
    """The column is written once; an amended file's newer shas live in the side table.

    This is plan Decision D2 in one test: the trigger stays and keeps meaning what it says, and
    ``latest_source_sha`` -- not the column -- is what a freeze check reads.
    """
    _method(conn)
    with conn:
        store.update_method(conn, "M0001", source_sha="a" * 64)
    with pytest.raises(store.LabError, match="set once"):
        store.update_method(conn, "M0001", source_sha="b" * 64)
    assert store.latest_source_sha(conn, "M0001") == "a" * 64
    with conn:
        store.insert_amendment(conn, store.AmendmentRow(
            method_id="M0001", seq=1, source_sha="c" * 64, added="2026-10-12",
            why="the brakes are the drag", hypothesis="h", expected_failure="e",
            run_at="2026-10-12T00:00:00+00:00", git_sha="deadbeef",
        ))
    assert store.get_method(conn, "M0001")["source_sha"] == "a" * 64
    assert store.latest_source_sha(conn, "M0001") == "c" * 64


def test_next_id_and_seen(conn):
    assert store.next_method_id(conn) == "M0001"
    _method(conn, "M0007")
    assert store.next_method_id(conn) == "M0008"
    with conn:
        assert store.mark_seen(conn, "url:https://example.com/x", "M0007", "a paper")
        assert not store.mark_seen(conn, "url:https://example.com/x")
    assert [r["key"] for r in store.seen(conn, "EXAMPLE")] == ["url:https://example.com/x"]
    with pytest.raises(store.LabError):
        store.mark_seen(conn, "no-kind")


def test_seed_imports_p7a_as_trials_1_to_54(conn):
    assert seed(conn) == 54
    rows = conn.execute("SELECT * FROM trials ORDER BY n").fetchall()
    assert [r["n"] for r in rows] == list(range(1, 55))
    assert [r["candidate_id"] for r in rows] == [c.id for c in REGISTRY[:54]]
    assert [r["config_digest"] for r in rows] == [config_digest(c) for c in REGISTRY[:54]]
    assert not any(r["eligible"] for r in rows)
    best = max(rows, key=lambda r: r["mar"] or -1)
    assert best["candidate_id"] == "F4-MOM12-N20-TREND"
    assert {r["status"] for r in conn.execute("SELECT status FROM methods")} == {"rejected"}
    with pytest.raises(store.LabError, match="runs once"):
        seed(conn)


def test_seed_configs_are_distinct():
    assert len({config_digest(c) for c in REGISTRY}) == len(REGISTRY)
    c = REGISTRY[3]
    assert config_digest(replace(c, id="F1-OTHER")) == config_digest(c)  # renaming is not a new trial


def test_export_writes_every_sheet(conn, tmp_path):
    from openpyxl import load_workbook

    seed(conn)
    path = store.export_xlsx(conn, tmp_path / "lab.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["summary", "leaderboard", "methods", "trials", "ideas_seen", "insights"]
    assert wb["leaderboard"].max_row == 55
    assert wb["leaderboard"]["B2"].value == "F4-MOM12-N20-TREND"


def test_insights_are_an_append_only_journal(conn):
    _method(conn)
    with conn:
        n = store.add_insight(conn, kind="data-wish", title="Fundamentals", body="Quarterly earnings would unlock value/quality.", method_id="M0001")
    assert n == 1
    with pytest.raises(store.LabError):
        store.add_insight(conn, kind="rumor", title="t", body="b")
    with pytest.raises(store.LabError):
        store.add_insight(conn, kind="risk", title="t", body="b", method_id="M0404")
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        conn.execute("UPDATE insights SET body = 'x'")
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        conn.execute("DELETE FROM insights")


def test_begin_immediate_makes_id_allocation_atomic(tmp_path):
    a = store.connect(tmp_path / "lab.sqlite")
    b = store.connect(tmp_path / "lab.sqlite")
    b.execute("PRAGMA busy_timeout = 0")
    store.begin_immediate(a)
    with pytest.raises(sqlite3.OperationalError, match="locked"):
        store.begin_immediate(b)
    a.rollback()
    store.begin_immediate(b)
    b.rollback()


# ---- promotion (roster-promotion-pipeline phase 5) ----------------------------------------------


def _promote(conn, **kw):
    base = dict(
        method_id="M0001", strategy_id="FND", candidate_id="M0001-A",
        object_name="FUNDAMENTAL", spec_digest="d" * 64,
        reason="the lab has not passed it; on paper to test it forward",
    )
    base.update(kw)
    with conn:
        return store.record_promotion(conn, **base)


def test_record_promotion_appends_analysis_and_an_insight(conn):
    _method(conn, status="idea")
    assert _promote(conn, move_status=False) == "idea"
    row = store.get_method(conn, "M0001")
    assert store.PROMOTION_MARKER + "`FND`" in row["analysis"]
    assert "M0001-A" in row["analysis"] and "FUNDAMENTAL" in row["analysis"]
    assert row["status"] == "idea"
    insight = conn.execute("SELECT * FROM insights ORDER BY id DESC LIMIT 1").fetchone()
    assert insight["kind"] == "observation" and insight["method_id"] == "M0001"
    assert "starts paper trading as FND" in insight["title"]
    assert "`" not in insight["body"] and "digest" not in insight["body"]


def test_record_promotion_takes_the_edge_transitions_already_has(conn):
    _method(conn, status="idea")
    with conn:
        for nxt in ("registered", "dev-eligible", "promoted", "test-passed"):
            store.update_method(conn, "M0001", status=nxt)
    assert _promote(conn) == "paper"
    assert store.get_method(conn, "M0001")["status"] == "paper"


def test_record_promotion_refuses_a_status_with_no_path_to_paper(conn):
    _method(conn, status="idea")
    with conn:
        store.update_method(conn, "M0001", status="rejected")
    with pytest.raises(store.LabError, match="--lab-status-stays"):
        _promote(conn)
    assert store.get_method(conn, "M0001")["status"] == "rejected"
    assert store.PROMOTION_MARKER not in store.get_method(conn, "M0001")["analysis"]


def test_record_promotion_is_idempotent_and_never_touches_source_sha(conn):
    _method(conn, status="idea")
    with conn:
        store.update_method(conn, "M0001", source_sha="a" * 64)
    _promote(conn, move_status=False)
    _promote(conn, move_status=False)
    row = store.get_method(conn, "M0001")
    assert row["analysis"].count(store.PROMOTION_MARKER) == 1
    assert conn.execute("SELECT count(*) FROM insights").fetchone()[0] == 1
    assert row["source_sha"] == "a" * 64


def test_promotion_basis_is_derived_from_the_lab_status():
    assert store.promotion_basis("test-passed") == "test-passed"
    assert store.promotion_basis("paper") == "test-passed"
    for s in ("idea", "registered", "rejected", "dev-eligible", "promoted", "blocked-data"):
        assert store.promotion_basis(s) == "owner-override"
    assert set(store.PROMOTION_BASES) == {"test-passed", "owner-override"}


def test_record_promotion_writes_the_basis_the_status_and_the_reason(conn):
    _method(conn, status="idea")
    assert _promote(conn, move_status=False) == "idea"
    analysis = store.get_method(conn, "M0001")["analysis"]
    assert "Lab status at admission: `idea`" in analysis
    assert "Basis: `owner-override`" in analysis
    assert "on paper to test it forward" in analysis


def test_an_override_with_no_reason_is_refused_and_writes_nothing(conn):
    _method(conn, status="idea")
    with pytest.raises(store.LabError, match="reason"):
        _promote(conn, move_status=False, reason="   ")
    row = store.get_method(conn, "M0001")
    assert store.PROMOTION_MARKER not in row["analysis"]
    assert conn.execute("SELECT count(*) FROM insights").fetchone()[0] == 0


def test_a_test_passed_method_needs_no_reason_and_says_so(conn):
    _method(conn, status="idea")
    with conn:
        for nxt in ("registered", "dev-eligible", "promoted", "test-passed"):
            store.update_method(conn, "M0001", status=nxt)
    assert _promote(conn, reason="") == "paper"
    analysis = store.get_method(conn, "M0001")["analysis"]
    assert "Basis: `test-passed`" in analysis
    assert "Lab status at admission: `test-passed`" in analysis


# ---- trial_provenance (schema v5, trial-reproducibility phase 2) ------------------------------


def _provenance(n: int, **kw) -> store.ProvenanceRow:
    base = dict(trial_n=n, initial_idr=Decimal("10000000"), price_fingerprint="pf",
                source="recorded", measured="2026-10-09T00:00:00+00:00")
    base.update(kw)
    return store.ProvenanceRow(**base)


def test_provenance_round_trips_the_capital_exactly_and_is_append_only(conn):
    _method(conn)
    with conn:
        (n,) = store.insert_trials(conn, [_trial()])
        store.insert_provenance(conn, [_provenance(n, initial_idr=Decimal("2E+7"))])
    row = store.provenance_of(conn, n)
    assert row["initial_idr"] == "20000000"  # format(d, "f"): exponent notation never stored
    assert Decimal(row["initial_idr"]) == Decimal("20000000")
    assert row["price_fingerprint"] == "pf" and row["source"] == "recorded"
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        conn.execute("UPDATE trial_provenance SET initial_idr = '1'")
    with pytest.raises(sqlite3.IntegrityError, match="append-only"):
        conn.execute("DELETE FROM trial_provenance")
    with pytest.raises(store.LabError, match="already has recorded provenance"):
        with conn:
            store.insert_provenance(conn, [_provenance(n)])


def test_provenance_refuses_what_a_rerun_could_not_use(conn):
    """A capital a run could not have started with, an unknown source, or the sentinel 0."""
    _method(conn)
    with conn:
        (n,) = store.insert_trials(conn, [_trial()])
    for bad, match in (
        (_provenance(0), "trial number"),
        (_provenance(n, initial_idr=Decimal("0")), "positive, finite"),
        (_provenance(n, initial_idr=Decimal("NaN")), "positive, finite"),
        (_provenance(n, initial_idr=10_000_000), "must be a Decimal"),
        (_provenance(n, source="guess"), "source"),
    ):
        with pytest.raises(store.LabError, match=match):
            with conn:
                store.insert_provenance(conn, [bad])
    assert store.provenance_of(conn, n) is None
    with conn:  # an unknown price fingerprint is recordable, and reads back as NULL
        store.insert_provenance(conn, [_provenance(n, price_fingerprint=None, source="backfill")])
    assert store.provenance_of(conn, n)["price_fingerprint"] is None


def test_seed_records_the_p7a_capital_and_price_fingerprint(conn):
    """The 58-trial cohort's 54 seed rows are marked, not rewritten (R4): 20M, the P7a prices."""
    seed(conn)
    rows = conn.execute("SELECT * FROM trial_provenance ORDER BY trial_n").fetchall()
    assert [r["trial_n"] for r in rows] == list(range(1, 55))
    assert {r["initial_idr"] for r in rows} == {"20000000"}
    assert {r["price_fingerprint"] for r in rows} == {seed_mod.P7A_FINGERPRINT}
    assert {r["source"] for r in rows} == {"backfill"}
    assert seed_mod.P7A_INITIAL_IDR == store.BACKFILL_LUMP_SUM_IDR
    assert seed_mod.P7A_FINGERPRINT == store.P7A_PRICE_FINGERPRINT
    assert store.BACKFILL_PRICE_FINGERPRINTS[seed_mod.P7A_FINGERPRINT] == store.P7A_PRICE_FINGERPRINT


# ----------------------------------------------------------- the duplicate note, end to end


def _note_db(tmp_path):
    db = tmp_path / "note.sqlite"
    c = store.connect(db)
    with c:
        store.add_method(
            c, id="M0001", name="placebo shape", family="control",
            source_kind="knowledge", hypothesis="h",
        )
    c.close()
    return db


def _lab(db, command, **kw):
    return lab_cmd.run(argparse.Namespace(db=db, lab_command=command, **kw))


def _read(db, mid="M0001"):
    c = store.connect(db)
    try:
        return dict(store.get_method(c, mid))
    finally:
        c.close()


def test_lab_note_says_what_it_wrote(tmp_path, capsys):
    """One line, naming the method and both halves of what landed. A pipe can match this."""
    db = _note_db(tmp_path)
    note = tmp_path / "note.md"
    note.write_text("The placebo is a control, not a constant.\n", encoding="utf-8")

    assert _lab(db, "note", method="M0001", file=note,
                verdict="  rejected -- the wrapper, not the signal  ") == 0
    out = capsys.readouterr().out.strip().splitlines()
    assert len(out) == 1
    assert out[0] == (
        'M0001: analysis appended (41 characters); '
        'verdict set to "rejected -- the wrapper, not the signal".'
    )
    row = _read(db)
    assert "The placebo is a control, not a constant." in row["analysis"]
    assert row["verdict"] == "rejected -- the wrapper, not the signal"


def test_lab_note_with_only_a_verdict_still_prints_one_line(tmp_path, capsys):
    db = _note_db(tmp_path)
    assert _lab(db, "note", method="M0001", file=None,
                verdict="rejected: the control was handicapped") == 0
    assert capsys.readouterr().out.strip() == (
        'M0001: verdict set to "rejected: the control was handicapped".'
    )
    row = _read(db)
    assert row["verdict"] == "rejected: the control was handicapped"
    assert row["analysis"] == ""


def test_lab_note_exits_2_on_the_second_identical_file_and_writes_nothing(tmp_path, capsys, caplog):
    """Exit 2 (the lab's rules refused), the refusal on the log, and the whole row untouched."""
    db = _note_db(tmp_path)
    note = tmp_path / "note.md"
    note.write_text("The placebo is a control, not a constant.\n", encoding="utf-8")
    assert _lab(db, "note", method="M0001", file=note, verdict="first verdict") == 0
    capsys.readouterr()
    before = _read(db)

    with caplog.at_level(logging.ERROR):
        assert _lab(db, "note", method="M0001", file=note, verdict="a different verdict") == 2
    assert "already the newest section" in caplog.text
    assert "lab show M0001" in caplog.text
    assert "no override" in caplog.text
    assert capsys.readouterr().out == ""          # nothing claims success

    after = _read(db)
    assert after["analysis"] == before["analysis"]
    assert after["verdict"] == "first verdict"    # the whole transaction rolled back


def test_the_duplicate_refusal_leaves_the_lab_s_own_generated_notes_alone(tmp_path, capsys):
    """``lab unblock`` and ``record_promotion`` both guard themselves before they ever append."""
    db = _note_db(tmp_path)
    assert _lab(db, "block", method="M0001", on="CBOE VIX3M daily closes") == 0
    assert _lab(db, "unblock", method="M0001", note="Covers 46.8% of the dev window.") == 0
    assert _lab(db, "block", method="M0001", on="EODHD split ex-dates") == 0
    assert _lab(db, "unblock", method="M0001", note="splits.csv arrived; the whole window.") == 0
    row = _read(db)
    assert row["status"] == "idea"
    assert row["analysis"].count("Unblocked: the missing data arrived.") == 2
    capsys.readouterr()

    c = store.connect(db)
    try:
        for _ in range(2):  # the second call returns on PROMOTION_MARKER, never reaching the append
            with c:
                store.record_promotion(
                    c, method_id="M0001", strategy_id="FND", candidate_id="M0001-A",
                    object_name="FUNDAMENTAL", spec_digest="d" * 64, move_status=False,
                    reason="the lab has not passed it; on paper to test it forward",
                )
        assert store.get_method(c, "M0001")["analysis"].count(store.PROMOTION_MARKER) == 1
    finally:
        c.close()
